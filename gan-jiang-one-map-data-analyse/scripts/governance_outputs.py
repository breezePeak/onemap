"""从逐页证据生成材料包、选址信息表、补建项目表及阶段映射，不修改原件。"""
import datetime as dt
import hashlib
import json
from pathlib import Path
import re
import shutil
import zipfile
import openpyxl
from pypdf import PdfReader,PdfWriter
from cross_stage_materials import read_json
from link_stages import target_records
from match_engine import normalize,rank_candidates


def safe_name(value):
    name=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',str(value or '未命名')).strip(' .')[:80] or '未命名'
    return '_'+name if re.fullmatch(r'CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9]',name,re.I) else name


def string_cell(sheet,row,column,value):
    cell=sheet.cell(row,column,value)
    if isinstance(value,str):cell.data_type='s'


def table(path,headers,rows):
    wb=openpyxl.Workbook();sheet=wb.active;sheet.title='明细'
    for i,record in enumerate([headers,*rows],1):
        for j,value in enumerate(record,1):string_cell(sheet,i,j,value)
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=24
    wb.save(path);wb.close()


def field_values(item,column):
    return list(dict.fromkeys(c['value'] for c in item.get('field_candidates',[]) if c['column']==column))


def single(item,column):
    values=field_values(item,column)
    return values[0] if len(values)==1 else None


def identity(item):
    text=item.get('text','')
    names=re.findall(r'(?:建设项目名称|项目名称|地块名称|工程名称)\s*[:：]\s*([^\n，,。；;:：]{2,100}?)(?=建设单位|项目代码|项目拟选位置|拟用地|\n|[，,。；;]|$)',text)
    fields=item.get('fields',{})
    return dict(name=names[0].strip() if len(set(names))==1 else None,
                **{k:values[0] for k,values in fields.items() if k in ('regulator_id','contract_id','unit_id','project_code') and len(values)==1})


def source_folder(source):
    label=source.split('#page=',1)[0].replace('\\','/').rsplit('!',1)[-1]
    return label.rsplit('/',2)[-2] if '/' in label else '未命名'


def rebuild_projects(items,targets):
    """合同/决定书明确在 2017 年后且无现有候选时补建；字段冲突保留待核实。"""
    projects={};deferred=[]
    for item in items:
        if item.get('projects') or item.get('pending_projects'):continue
        ident=identity(item);date=single(item,68)
        if any(len(item.get('fields',{}).get(k,[]))>1 for k in ('regulator_id','contract_id','unit_id','project_code')):
            deferred.append(dict(source=item['source'],reason='同页多个标识，需先拆分或确认关系',identity=ident));continue
        stage=item.get('material_stage')
        number=single(item,67 if stage=='土地供应' else 43 if stage=='规划条件与红线' else 35 if stage=='预审与选址' else 52)
        if stage not in ('土地供应','规划条件与红线','预审与选址','用地报批') or not ident.get('name') or not (ident.get('contract_id') or number):
            deferred.append(dict(source=item['source'],reason='缺明确业务文书、项目名称或文书编号',identity=ident));continue
        if stage=='土地供应':
            try:recent=dt.date(2018,1,1)<=dt.date.fromisoformat(str(date))<=dt.date.today()
            except ValueError:recent=False
            if not recent:
                deferred.append(dict(source=item['source'],reason='合同/决定书日期缺失、冲突或不在2017年后',identity=ident));continue
        if item.get('identity_conflicts') or any(any(normalize(ident.get(k)) and normalize(ident[k])==normalize(t.get(k)) for k in ('contract_id','regulator_id','unit_id')) for t in targets):
            deferred.append(dict(source=item['source'],reason='与已有项目存在相同编号或标识冲突，需消歧',identity=ident));continue
        key=(normalize(ident.get('regulator_id')),normalize(ident.get('contract_id') or stage+':'+str(number)))
        project=projects.setdefault(key,dict(identity=ident,archive_id=source_folder(item['source']),values={},sources=[],status='待核实补建项目，项目类型待标准核定'))
        for c in item.get('field_candidates',[]):
            values=project['values'].setdefault(str(c['column']),[])
            if c['value'] not in values:values.append(c['value'])
            if len(values)>1:project['status']='同编号文书字段冲突，需人工消歧'
        for k,v in ident.items():
            if project['identity'].get(k) and normalize(project['identity'][k])!=normalize(v):project['status']='同编号文书字段冲突，需人工消歧'
        project['sources'].append(item['source'])
    complete=[]
    for project in projects.values():
        if project['status']=='同编号文书字段冲突，需人工消歧':deferred.append(dict(sources=project['sources'],reason=project['status'],identity=project['identity']))
        else:complete.append(project)
    return complete,deferred


def registration_projects(job,items,targets):
    projects=[];pending=[]
    records=read_json(job/'work/地表宗地.json',[])
    counts={}
    for record in records:
        unit=normalize(record.get('BDCDYH'));counts[unit]=counts.get(unit,0)+1
    for index,record in enumerate(records):
        unit=record.get('BDCDYH')
        if not normalize(unit):continue
        candidates=rank_candidates(dict(unit_id=unit,name=record.get('XMMC'),address=record.get('ZL')),targets)
        document_hits=[i['source'] for i in items if unit in i.get('fields',{}).get('unit_id',[]) and (i.get('projects') or i.get('pending_projects'))]
        if candidates or document_hits or counts[normalize(unit)]>1:
            pending.append(dict(unit_id=unit,source_record=index,project_candidates=[c['record']['project'] for c in candidates],document_candidates=document_hits,
                                reason='同单元号多条图层记录，保留全部关系待核对' if counts[normalize(unit)]>1 else '存在已有项目或文书候选，待消歧'));continue
        projects.append(dict(identity=dict(name=record.get('XMMC'),unit_id=unit),archive_id='',values={},
                             original_values={12:record.get('ZL'),14:record.get('ZDDM')},sources=[str(job/'work/地表宗地.json')+'#record='+str(index)],
                             status='依据首登宗地图层补建；项目名称、权利关系和项目类型待核实'))
    return projects,pending


def build(job):
    job=Path(job);out=job/'outputs/governance';out.mkdir(parents=True,exist_ok=True)
    evidence=out/'识别字段候选.jsonl'
    items=[json.loads(line) for line in evidence.read_text(encoding='utf-8').splitlines() if line.strip()] if evidence.exists() else []
    # 只在同页同一实体内合并识别方式，不能把已拆分的项目再次拼成一行。
    pages={}
    for item in items:
        page_key=(item['source'],item.get('block_id','block-1'))
        if page_key not in pages:pages[page_key]=dict(item,field_candidates=list(item.get('field_candidates',[])),fields={k:list(v) for k,v in item.get('fields',{}).items()})
        else:
            old=pages[page_key]
            for c in item.get('field_candidates',[]):
                if c not in old['field_candidates']:old['field_candidates'].append(c)
            for key,values in item.get('fields',{}).items():old['fields'][key]=list(dict.fromkeys(old['fields'].get(key,[])+values))
            if item.get('text') and item['text']!=old.get('text'):old['text']=old.get('text','')+'\n'+item['text']
            if old.get('material_stage','未分类')=='未分类':old['material_stage']=item.get('material_stage','未分类')
            elif item.get('material_stage','未分类') not in ('未分类',old['material_stage']):old['material_stage']='未分类';old['stage_conflict']=True
            for key in ('projects','pending_projects','identity_conflicts'):
                old[key]=sorted(set(old.get(key,[]))|set(item.get(key,[])))
            if len(old.get('projects',[]))>1 or old.get('identity_conflicts'):
                old['pending_projects']=sorted(set(old.get('pending_projects',[]))|set(old.get('projects',[])))
                old['projects']=[]
    items=list(pages.values());targets=target_records(job/'test_result.xlsx')
    supply_folders={}
    for item in [*items,*read_json(job/'work/pdf_documents.json',[]),*read_json(job/'work/ocr_manifest.json',[])]:
        if '土地供应' not in str(item.get('storage_stage')):continue
        parts=item['source'].replace('\\','/').replace('!','/').split('/')
        for i,part in enumerate(parts[:-1]):
            if re.fullmatch(r'20\d{2}-\d{3,}',part) and dt.date.today().year-9<=int(part[:4])<=dt.date.today().year:
                folder='/'.join(parts[:i+1]);entry=supply_folders.setdefault(folder,dict(archive=part,projects=set(),sources=set()))
                entry['projects'].update(item.get('projects',[])+item.get('pending_projects',[])+item.get('retrieval_projects',[]));entry['sources'].add(item['source'].split('#page=',1)[0]);break
    table(out/'供应目录建项清单.xlsx',['来源目录','文件夹编号','材料数','现有项目候选','状态'],
          [[folder,e['archive'],len(e['sources']),'；'.join(sorted(e['projects'])),'已有项目候选，待核对' if e['projects'] else '目录待建项；核对文书日期、名称和编号后进入补建表'] for folder,e in supply_folders.items()])
    documents={d['source']:d for d in read_json(job/'work/pdf_documents.json',[])}
    material_rows=[];selection_rows=[];errors=[];mappings=[]
    for item in items:
        stage=item.get('material_stage','未分类')
        if item.get('stage_conflict'):
            errors.append(dict(source=item['source'],error='同页不同识别方式的文书类型冲突，待人工核对'));continue
        if stage=='未分类':continue
        source=item['source'];label=source.split('#page=',1)[0]
        number=single(item,43) or single(item,35) or single(item,52) or single(item,67)
        digest=hashlib.sha256(source.encode()).hexdigest()[:16]
        folder=out/'阶段材料'/safe_name(stage)/safe_name(number or '待补编号')
        destination=None
        try:
            folder.mkdir(parents=True,exist_ok=True)
            document=documents.get(label)
            if document and '#page=' in source:
                destination=folder/(safe_name(source_folder(source))+'_'+safe_name('用地预审与选址意见书' if stage=='预审与选址' else stage)+'_'+digest+'.pdf')
                with Path(document['path']).open('rb') as stream:
                    reader=PdfReader(stream);writer=PdfWriter();writer.add_page(reader.pages[int(source.rsplit('#page=',1)[1])-1])
                    temp=destination.with_suffix('.tmp');writer.write(temp);writer.close();temp.replace(destination)
            else:
                image=Path(item.get('image') or label)
                if image.is_file():
                    destination=folder/(safe_name(source_folder(source))+'_'+safe_name(stage)+'_'+digest+image.suffix)
                    shutil.copy2(image,destination)
                else:raise FileNotFoundError('缺原件/页面图像，仅保留文字提取记录')
        except Exception as error:errors.append(dict(source=source,error=str(error)));destination=None
        ident=identity(item);fields=item.get('fields',{})
        status='页内归属已定位，字段待核对' if item.get('projects') else '项目归属待核对'
        material_rows.append([str(destination.relative_to(out)) if destination else None,stage,number,source,item.get('storage_stage'),';'.join(item.get('projects',[])), ';'.join(item.get('pending_projects',[])),status])
        if stage=='预审与选址':
            selection_rows.append([str(destination.relative_to(out)) if destination else source,ident.get('name'),ident.get('project_code'),
                                   *['；'.join(map(str,field_values(item,col))) for col in (30,32,33,34,35,36,37)],source,status])
        for regulator in fields.get('regulator_id',[]):
            for unit in fields.get('unit_id',[]):
                mappings.append([regulator,unit,'；'.join(fields.get('contract_id',[])),source,'文书同页共现，待确认' if len(fields.get('regulator_id',[]))==len(fields.get('unit_id',[]))==1 else '同页多编号，需消歧'])
    projects,deferred=rebuild_projects(items,targets)
    registered,registration_pending=registration_projects(job,items,targets)
    document_units={normalize(p['identity'].get('unit_id')) for p in projects if p['identity'].get('unit_id')}
    projects.extend(p for p in registered if normalize(p['identity']['unit_id']) not in document_units)
    table(out/'首登补建待消歧.xlsx',['单元号','图层记录位置','已有项目候选','文书候选','原因'],[[p['unit_id'],p['source_record'],'；'.join(p['project_candidates']),'；'.join(p['document_candidates']),p['reason']] for p in registration_pending])
    for asset in read_json(job/'work/material_assets.json',[]):
        source=asset['source'];path=Path(asset['path']);stage='规划条件与红线' if '红线' in source else '待分类CAD附件'
        folder=out/'阶段材料'/stage/'待补编号';folder.mkdir(parents=True,exist_ok=True)
        destination=folder/(safe_name(path.stem)+'_'+hashlib.sha256(source.encode()).hexdigest()[:12]+path.suffix)
        try:
            shutil.copy2(path,destination)
            material_rows.append([str(destination.relative_to(out)),stage,None,source,asset.get('storage_stage'),None,'；'.join(asset.get('retrieval_projects',[])),'CAD原件已整理，图形及项目归属待判读'])
        except OSError as error:errors.append(dict(source=source,error=str(error)))
    for index,project in enumerate(projects,1):project['id']='new-'+str(index)
    base=out/'test_result_已填充.xlsx'
    wb=openpyxl.load_workbook(base if base.exists() else job/'test_result.xlsx');sheet=wb.worksheets[0]
    for project in projects:
        row=sheet.max_row+1;ident=project['identity'];values={2:project['archive_id'],24:ident.get('name'),15:ident.get('unit_id'),18:ident.get('regulator_id'),19:ident.get('contract_id')}
        values.update(project.get('original_values',{}))
        values.update({int(k):v[0] for k,v in project['values'].items() if len(v)==1})
        for col,value in values.items():
            if value is not None:string_cell(sheet,row,col,value)
        project['row']=row
    audit=wb.create_sheet('补建依据' if '补建依据' not in wb.sheetnames else '本次补建依据')
    for i,record in enumerate([['新增编号','目标行','来源','状态'],*[[p['id'],p['row'],'；'.join(p['sources']),p['status']] for p in projects]],1):
        for j,value in enumerate(record,1):string_cell(audit,i,j,value)
    wb.save(out/'项目表_含补建项目.xlsx');wb.close()
    (out/'补建项目.json').write_text(json.dumps(dict(projects=projects,deferred=deferred),ensure_ascii=False,indent=2),encoding='utf-8')
    types=[]
    for spatial in read_json(job/'work/spatial_matches.json',[]):
        results=[c['result'] for c in spatial['candidates']] if spatial.get('supply_pending') else [spatial]
        for result in results:
            if result.get('resolution',{}).get('excluded'):continue
            for hit in result.get('approvals',[]):
                if hit.get('manual_verdict')=='不一致':continue
                attrs=hit['attributes']
                if attrs.get('XMLX'):types.append([spatial['project'],attrs.get('OBJECTID'),attrs['XMLX'],hit['supply_coverage'],'项目类型新标准未提供，待核定，不直接入库'])
            for hit in result.get('parcels',[]):
                if hit.get('manual_verdict')=='不一致':continue
                unit=hit['attributes'].get('BDCDYH')
                target=next((t for t in targets if t['project']==spatial['project']),{})
                if unit and target.get('regulator_id'):mappings.append([target['regulator_id'],unit,target.get('contract_id'),str(job/'work/spatial_matches.json'),'空间相交候选，需核对覆盖率及文书'])
    table(out/'选址意见书信息提取.xlsx',['文件名称','项目名称','项目代码','项目拟选位置','拟用地面积','拟建设规模','项目建设依据','证书文号','发证日期','面积单位','原始来源','核对状态'],selection_rows)
    table(out/'阶段材料清单.xlsx',['整理后文件','业务阶段','文书编号','原始来源','存放阶段','已定位项目','待确认项目','状态'],material_rows)
    table(out/'监管号与单元号映射.xlsx',['电子监管号','不动产单元号','合同号','来源','状态'],mappings)
    table(out/'项目类型核定清单.xlsx',['项目','报批图斑OBJECTID','原始项目类型','参考范围覆盖率','状态'],types)
    table(out/'补建项目清单.xlsx',['新增编号','项目名称','目标行','来源档案目录','监管号','合同/决定书编号','来源','状态'],
          [[p['id'],p['identity']['name'],p['row'],p['archive_id'],p['identity'].get('regulator_id'),p['identity'].get('contract_id') or next(iter(p['values'].get('67',[])),None),'；'.join(p['sources']),p['status']] for p in projects])
    source_errors=[]
    for filename,stage in [('material_errors.json','供应扫描件'),('source_read_errors.json','来源读取'),('pdf_errors.json','材料解析'),('pdf_render_errors.json','PDF渲染'),('ocr_errors.json','OCR')]:
        for error in read_json(job/'work'/filename,[]):source_errors.append(dict(stage=stage,**error))
    source_errors.extend(dict(stage='附件矢量',**e) for e in read_json(job/'outputs/governance/矢量整理统计.json',{}).get('errors',[]))
    table(out/'来源异常清单.xlsx',['处理环节','原始来源','页面路径','异常原因'],[[e['stage'],e.get('source'),e.get('path'),e.get('error')] for e in source_errors])
    (out/'来源异常清单.json').write_text(json.dumps(source_errors,ensure_ascii=False,indent=2),encoding='utf-8')
    report=dict(recognized_pages=len({i['source'] for i in items}),recognized_entities=len(items),source_errors=len(source_errors),organized_materials=sum(bool(r[0]) for r in material_rows),selection_records=len(selection_rows),new_projects=len(projects),deferred_projects=len(deferred),
                material_errors=errors,standards_pending=['目录组织标准未提供，当前仅按阶段和文书编号整理','项目类型新标准未提供，补建表不代表已审核入库'])
    (out/'成果生成统计.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    with zipfile.ZipFile(out/'治理成果.zip','w',zipfile.ZIP_DEFLATED) as bundle:
        for path in out.rglob('*'):
            if path.is_file() and path.name!='治理成果.zip' and not path.name.endswith('.tmp'):bundle.write(path,path.relative_to(out))
    return report


if __name__=='__main__':print(json.dumps(build(Path.cwd()),ensure_ascii=False))
