"""生成空白单元格补丁和证据清单。冲突保留原值，禁止无证据补齐。"""
import csv
import json
import re
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter
from inspect_sources import OUT,ROOT
from match_engine import normalize
from link_stages import is_project_row

DEST=Path('outputs/governance')

def main():
    DEST.mkdir(parents=True,exist_ok=True)
    w=openpyxl.load_workbook('test_result.xlsx',read_only=True,data_only=False);s=w.worksheets[0]
    rows={f'row-{i}':(i,list(r)) for i,r in enumerate(s.iter_rows(min_row=3,max_col=69,values_only=True),3) if is_project_row(r)}
    headers=[c.value for c in next(s.iter_rows(min_row=2,max_row=2,max_col=69))]
    w.close()
    patches={}; conflicts=[]
    def offer(project,column,value,source,method,**extra):
        if value is None or value=='':return
        if '日期' in str(headers[column-1]):
            assert extra.get('type')=='date' and re.fullmatch(r'\d{4}-\d{2}-\d{2}',str(value)),f'日期字段禁止写入非日期值: {headers[column-1]}'
        row,values=rows[project];cell=f'{get_column_letter(column)}{row}'
        old=values[column-1]
        if old is not None and str(old).strip():
            if normalize(old)!=normalize(value):conflicts.append(dict(project=project,cell=cell,field=headers[column-1],existing=old,candidate=value,source=source,reason='已有值与来源不同，保留原值'))
            return
        patch=dict(project=project,cell=cell,field=headers[column-1],value=value,source=source,method=method,**extra)
        if cell in patches and normalize(patches[cell]['value'])!=normalize(value):
            conflicts.append(dict(project=project,cell=cell,existing=patches[cell]['value'],candidate=value,source=source,reason='多个候选值冲突，撤销回填'));patches[cell]['blocked']=True
        elif cell not in patches:patches[cell]=patch
    # 只使用原始 XMMC/XMDM 配对，重复项目名或不同项目代码禁止独断。
    ledger=json.loads((OUT/'6.用地规划_台账.json').read_text(encoding='utf-8'))
    for key,(row,values) in rows.items():
        matches=[(i,r) for i,r in enumerate(ledger[1:],2) if normalize(r[12])==normalize(values[23])]
        codes={normalize(r[9]) for _,r in matches if r[9]}
        if len(codes)==1:
            if values[24] and normalize(values[24]) not in codes:continue
            for i,r in matches:
                offer(key,25,r[9],f'6.用地规划原始台账!J{i},M{i}','台账挂接：项目名称与台账XMMC唯一一致，回填该行项目代码(XMDM)')
                offer(key,26,r[18],f'6.用地规划原始台账!S{i}','台账挂接：项目名称与XMMC、项目代码与XMDM一致，回填建设单位')
    # 强空间证据：唯一批次覆盖供地至少 98%，其他批次均不足 1%。
    # 只回填批复文号/日期，不把批次名称写进“报批日期”或复制整批面积。
    for result in json.loads((OUT/'spatial_matches.json').read_text(encoding='utf-8')):
        if result.get('supply_pending') or result.get('anchor_layer','supply')!='supply' or not result.get('resolution',{}).get('automatic_allowed',True):continue
        values=rows[result['project']][1]
        declared=values[60]
        if isinstance(declared,(int,float)) and values[61]=='平方米' and abs(result['area']-declared)/declared>.02:
            continue  # 图层可能只含项目的一部分，不能据此认定项目只有该批次。
        grouped={}
        for hit in result['approvals']:
            if hit.get('manual_verdict')=='不一致':continue
            a=hit['attributes'];doc=a.get('PFWH') or a.get('PZWH')
            if not doc:continue
            group=grouped.setdefault(doc,{'coverage':0,'hits':[]})
            coverage=next((g['coverage'] for g in result.get('groups',[]) if g['layer']=='approvals' and g['key']==str(doc)),hit['supply_coverage'])
            group['coverage']=max(group['coverage'],coverage);group['hits'].append(a)
        strong=[(doc,g) for doc,g in grouped.items() if g['coverage']>=.98]
        if len(strong)!=1 or any(g['coverage']>=.01 for doc,g in grouped.items() if doc!=strong[0][0]):continue
        doc,g=strong[0];a=g['hits'][0]
        date=a.get('PZSJ')
        source='批文报批矢量1016.dbf OBJECTID='+','.join(x['OBJECTID'] for x in g['hits'])
        offer(result['project'],52,doc,source,'空间挂接：唯一文号批次累计覆盖供地≥98%且其他批次均<1%，回填批复文号(PFWH/PZWH)')
        if date and re.fullmatch(r'\d{8}',date) and date>'19000101':
            offer(result['project'],53,date[:4]+'-'+date[4:6]+'-'+date[6:],source,'空间挂接：同一批次批复日期(PZSJ)',type='date')
    review=Path('reviewed_fields.json')
    if review.exists():
        for item in json.loads(review.read_text(encoding='utf-8')):
            matches=[key for key,(_,values) in rows.items() if str(values[1])==str(item['project'])]
            if len(matches)!=1:continue
            assert Path(item['evidence_image']).is_file(),'人工核验图像不存在'
            offer(matches[0],item['column'],item['value'],item['source'],'人工逐页视觉核验（reviewed_fields.json，附证据图像）',evidence_image=item['evidence_image'])
    patches=[x for x in patches.values() if not x.get('blocked')]
    (OUT/'patches.json').write_text(json.dumps(patches,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    (OUT/'search_patches.json').write_text(json.dumps(patches,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    def csv_write(name,records,fields):
        with (DEST/name).open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(records)
    csv_write('填充溯源.csv',patches,['project','cell','field','value','method','source','evidence_image'])
    csv_write('已有值冲突.csv',conflicts,['project','cell','field','existing','candidate','source','reason'])
    changed={x['cell'] for x in patches};missing=[]
    for key,(row,values) in rows.items():
        for column in range(24,70):
            cell=f'{get_column_letter(column)}{row}'
            if (values[column-1] is None or not str(values[column-1]).strip()) and cell not in changed:
                if column in (28,29):reason='特征码编码规则未给出，不能推测'
                elif column in (25,26):reason='未取得本项目唯一代码/建设单位证据，参见挂接日志及识别队列'
                elif 30<=column<=38:reason='需核实本项目预审/选址材料；可从其他阶段档案取原件，不套用其他阶段字段值'
                elif 39<=column<=46:reason='缺本项目规划条件证据或图像需核验；不把选址红线当作规划条件'
                elif 47<=column<=55:reason='批次一对多、空间覆盖不足或缺单项目面积/耕地等法定属性'
                else:reason='现有来源无确定值，参见待复核队列'
                missing.append(dict(project=key,cell=cell,field=headers[column-1],status='PENDING',reason=reason))
    csv_write('剩余空值.csv',missing,['project','cell','field','status','reason'])
    summary=dict(projects=len(rows),filled_cells=len(patches),remaining_business_blanks=len(missing),existing_value_conflicts=len(conflicts))
    (DEST/'填充统计.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(summary,ensure_ascii=False))

if __name__=='__main__':main()
