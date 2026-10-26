"""本地工作台的检索结果、字段候选、选择记录与导出。"""
import csv
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import time
import zipfile
import openpyxl
from openpyxl.utils import get_column_letter
from match_engine import decide,normalize
from link_stages import is_project_row,row_identity,target_records

BASE=Path(__file__).resolve().parent.parent
from source_config import SOURCE
JOBS=BASE/'work'/'jobs'
RUNTIME=Path(os.environ['USERPROFILE'])/'.cache/codex-runtimes/codex-primary-runtime/dependencies'
NODE=RUNTIME/'node/bin/node.exe'
HEADERS=list(next(openpyxl.load_workbook(BASE/'test_result.xlsx',read_only=True).worksheets[0].iter_rows(min_row=2,max_row=2,max_col=69,values_only=True)))

def read_json(path,default=None):
    return json.loads(path.read_text(encoding='utf-8')) if path.exists() else default

def write_json(path,value):
    path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix(path.suffix+'.tmp');temp.write_text(json.dumps(value,ensure_ascii=False,default=str),encoding='utf-8')
    for attempt in range(4):
        try:
            temp.replace(path)
            return
        except PermissionError as error:
            if os.name!='nt' or getattr(error,'winerror',None) not in (5,32,33) or attempt==3: raise
            time.sleep(0.05*(attempt+1))

def value_json(value):
    return value.isoformat()[:10] if isinstance(value,(dt.date,dt.datetime)) else value

def value_key(value):
    if isinstance(value,(float,int)) and not isinstance(value,bool):return 'number:'+str(float(value))
    return 'text:'+normalize(value)

def validate_upload(path):
    if not zipfile.is_zipfile(path):raise ValueError('请上传真正的 .xlsx 文件，旧版 .xls 需先另存为 .xlsx。')
    with zipfile.ZipFile(path) as z:
        if sum(x.file_size for x in z.infolist())>100*1024*1024 or len(z.infolist())>2000:raise ValueError('工作簿解压后超过大小限制。')
    w=openpyxl.load_workbook(path,read_only=True,data_only=False)
    try:
        s=w.worksheets[0]

        headers=list(next(s.iter_rows(min_row=2,max_row=2,max_col=69,values_only=True)))
        if [normalize(x) for x in headers]!=[normalize(x) for x in HEADERS]:raise ValueError('列结构与当前 69 列治理模板不一致，请使用当前模板。')
        if s.max_row>5000:raise ValueError('首版单次最多处理 4998 个项目。')
        if not any(is_project_row(r) for r in s.iter_rows(min_row=3,max_col=69,values_only=True)):
            raise ValueError('未找到项目，请检查第 3 行起的项目名称、来源编号或宗地标识。')
    finally:w.close()

def run_step(job,script,*args):
    env=os.environ.copy();env['PYTHONIOENCODING']='utf-8';env['GOV_ASSETS']=str(BASE/'work')
    env['PYTHONPATH']=str(BASE/'work/python_packages')+os.pathsep+env.get('PYTHONPATH','')
    executable=str(NODE) if script.endswith('.mjs') else 'powershell' if script.endswith('.ps1') else sys.executable
    command=[executable]
    if script.endswith('.ps1'):command+=['-NoProfile','-ExecutionPolicy','Bypass','-File']
    command+=[str(BASE/'scripts'/script),*args]
    with (job/'console.log').open('a',encoding='utf-8') as log:
        log.write('\n>>> '+script+'\n');log.flush()
        p=subprocess.run(command,cwd=job,env=env,stdout=log,stderr=subprocess.STDOUT,timeout=3600,creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
    if p.returncode:raise RuntimeError(script+' 失败，查看运行日志；任务未被标记完成。')

def prepare_job(job):
    (job/'work').mkdir(parents=True,exist_ok=True)
    reviews=read_json(BASE/'reviewed_fields.json',[])
    for r in reviews:r['evidence_image']=str((BASE/r['evidence_image']).resolve())
    write_json(job/'reviewed_fields.json',reviews)
    for name in ('hash_cache.json',):
        if (BASE/'work'/name).exists():shutil.copy2(BASE/'work'/name,job/'work'/name)
    reuse_material_cache(job,BASE/'work')


def reuse_material_cache(job,assets):
    """只继承识别产物；项目归属、字段匹配和用户选择由当前任务重算。"""
    destination=job/'work';destination.mkdir(parents=True,exist_ok=True)
    caches=read_json(destination/'pdf_text_cache.json',{})
    roots=[assets,assets/'cross_stage_validation'/'work',*sorted(assets.glob('jobs/*/work'))]
    copied=set()
    for root in roots:
        if root.resolve()==destination.resolve():continue
        for key,record in read_json(root/'pdf_text_cache.json',{}).items():
            if record.get('error'):continue
            path=Path(key)
            if not path.is_file():continue
            stat=path.stat()
            if record.get('size')==stat.st_size and record.get('mtime_ns')==stat.st_mtime_ns:
                caches[key]=record
        for fragment in (root/'render_manifests').glob('*.json'):
            record=read_json(fragment,{})
            if not record.get('complete') or record.get('error'):continue
            if not all(Path(p['path']).is_file() for p in record.get('pages',[])):continue
            target=destination/'render_manifests'/fragment.name
            current=read_json(target,{})
            if current.get('mtimeMs',-1)>record.get('mtimeMs',-1):continue
            if target.resolve()==fragment.resolve():continue
            write_json(target,record);copied.add(record['source'])
    write_json(destination/'pdf_text_cache.json',caches)
    write_json(job/'cache_reuse.json',dict(pdf_text_documents=len(caches),render_records=len(copied),
                                         note='仅继承成功产物；渲染器继续核对源文件版本，归属与字段重新分析'))

def collect_review(job):
    previous=read_json(job/'review.json',{})
    identities={t['row']:t for t in target_records(job/'test_result.xlsx')}
    w=openpyxl.load_workbook(job/'test_result.xlsx',read_only=True,data_only=False)
    sheet=w.worksheets[0];rows=[];cells={};sources={}
    files=read_json(job/'work/files.json',[])
    by_name={Path(x['path']).name:Path(x['path']) for x in files}
    documents=read_json(job/'work/pdf_documents.json',[])
    manifest=read_json(job/'work/ocr_manifest.json',[])
    def source_ref(label,location='',image=None,detail=None):
        label=str(label);path=None
        location=location.removeprefix('原始台账!')
        if '.dbf OBJECTID=' in label:location='OBJECTID='+label.split('OBJECTID=',1)[1]
        if label.startswith(('6.用地规划原始台账!','7.工程规划原始台账!')):
            branch=label.split('原始台账!')[0];location=label.split('!',1)[1]
            path=SOURCE/branch/'ZRZYSITE_CONSTRUCT__001012_003b.csv';label=str(path)
        if image and Path(image).is_file():path=Path(image).resolve()
        if path is None:
            possible=Path(label)
            if possible.is_file():path=possible.resolve()
        if path is None:
            document=next((d for d in documents if label.split('#page=',1)[0]==d['source'] or ('.pdf' in label and label.split('!')[-1].split('，')[0] in d['source'])),None)
            if document:path=Path(document['path']).resolve()
        if path is None:
            match=next((p for name,p in by_name.items() if name in label and p.suffix.lower() in ('.dbf','.xlsx')),None)
            if match:path=match.resolve()
        if path and not any(path.is_relative_to(root.resolve()) for root in (BASE,SOURCE)):path=None
        ident=hashlib.sha256(((str(path) if path else label)+'|'+location).encode()).hexdigest()[:20]
        sources.setdefault(ident,dict(id=ident,label=label,location=location,path=str(path) if path else None,detail=detail,kind=path.suffix.lower().lstrip('.') if path else 'text'))
        if detail:sources[ident]['detail']=detail
        return ident
    def add(cell,value,source,method,auto=False,value_type=None,project_pending=False):
        if value is None or value=='':return
        entry=cells[cell];value=value_json(value)
        if value_type=='date':
            try:value=dt.datetime.fromisoformat(str(value)).date().isoformat()
            except ValueError:return
        key=value_key(value)
        candidate=next((c for c in entry['candidates'] if c['key']==key),None)
        if candidate is None:
            candidate=dict(id=hashlib.sha256((cell+'|'+key).encode()).hexdigest()[:16],key=key,value=value,type=value_type,sources=[],methods=[],auto=False,project_pending=project_pending)
            entry['candidates'].append(candidate)
        if source not in candidate['sources']:candidate['sources'].append(source)
        if method not in candidate['methods']:candidate['methods'].append(method)
        candidate['auto']=candidate['auto'] or auto
        candidate['project_pending']=candidate.get('project_pending',False) and project_pending
    for rownum,record in enumerate(sheet.iter_rows(min_row=3,max_col=69,values_only=True),3):
        if not is_project_row(record):continue
        row=dict(row=rownum,project=f'row-{rownum}',archive_id=str(record[1] or ''),name=str(record[23] or record[9] or record[2] or ''),cells=[])
        for col,value in enumerate(record,1):
            if col<24:continue
            cell=f'{get_column_letter(col)}{rownum}';row['cells'].append(cell)
            cells[cell]=dict(cell=cell,column=col,field=HEADERS[col-1],project=row['project'],original=value_json(value),value=value_json(value),candidates=[],status='original' if value is not None else 'empty')
        rows.append(row)
    w.close()
    targets={r['project']:r for r in rows}
    search_patches=read_json(job/'work/search_patches.json')
    if search_patches is None:
        search_patches=read_json(job/'work/patches.json',[])
        write_json(job/'work/search_patches.json',search_patches)
    for p in search_patches:
        if p['cell'] in cells:add(p['cell'],p['value'],source_ref(p['source'],image=p.get('evidence_image')),p['method'],True,p.get('type'))
    # 成果总表逐字段挂接：精确标识无冲突的唯一项目才自动通过。
    master=SOURCE/'project_info_all_v5_a - 成果20261021.xlsx'
    master_records=[]
    if master.is_file():
        mw=openpyxl.load_workbook(master,read_only=True,data_only=True)
        try:
            master_title=mw.worksheets[0].title
            for i,r in enumerate(mw.worksheets[0].iter_rows(min_row=3,max_col=69,values_only=True),3):
                if not any(r):continue
                master_records.append(dict(row_identity(r),values=list(r),source_row=i))
        finally:mw.close()
    tw=openpyxl.load_workbook(job/'test_result.xlsx',read_only=True,data_only=True)
    key_labels=dict(project_code='项目代码',regulator_id='电子监管号',contract_id='合同号',unit_id='不动产单元号',condition_id='规划条件出具编号',approval_id='批复文号',certificate_id='选址证书文号')
    for i,r in enumerate(tw.worksheets[0].iter_rows(min_row=3,max_col=69,values_only=True),3):
        if not is_project_row(r):continue
        target=identities[i]
        status,ranked=decide(target,master_records)
        for match in ranked[:5]:
            if match['conflicts']:continue
            rec=match['record']
            basis='、'.join(key_labels.get(k,k) for k in match['exact'])
            if status=='MATCH_EXACT':method='成果总表：'+basis+'与本项目完全一致，唯一且无冲突（精确挂接）'
            elif match['exact']:method='成果总表：'+basis+'一致，但存在多条记录或依据为尚未确认的识别编号，需人工确认'
            else:method='成果总表：名称相似度{:.0%}、地址相似度{:.0%}、面积相对误差{}，综合候选需人工确认'.format(match['name_score'],match['address_score'],'未知' if match['area_error'] is None else '{:.1%}'.format(match['area_error']))
            for col,value in enumerate(rec['values'],1):
                if col<24 or value is None:continue
                cell=f'{get_column_letter(col)}{i}'
                ref=source_ref(str(master),f'{master_title}!{get_column_letter(col)}{rec["source_row"]}')
                add(cell,value,ref,method,status=='MATCH_EXACT','date' if isinstance(value,(dt.date,dt.datetime)) else None)
    tw.close()
    # 目录线索可展示，但必须明确确认项目归属后才允许采用。
    evidence_file=job/'outputs/governance/识别字段候选.jsonl'
    if evidence_file.exists():
        for line in evidence_file.read_text(encoding='utf-8').splitlines():
            if not line.strip():continue
            item=json.loads(line)
            assigned=item.get('projects',[]);pending=not assigned
            projects=assigned or item.get('pending_projects',[])
            for project in projects:
                row=targets.get(project)
                if not row:continue
                for candidate in item.get('field_candidates',[]):
                    column=candidate['column']
                    if not 24<=column<=69:continue
                    page='第'+item['source'].split('#page=')[-1]+'页' if '#page=' in item['source'] else '图片页'
                    location=page+'，'+(item['block_id']+'，' if item.get('block_id') else '')+get_column_letter(column)+'列，'+item['text_method']
                    ref=source_ref(item['source'],location,image=item.get('image'),
                                   detail=dict(实体分段=item.get('block_id'),分段依据=item.get('segmentation'),OCR修正候选=item.get('ocr_corrections',[]),存放阶段=item.get('storage_stage'),材料所属阶段=candidate['material_stage'],识别方式=item['text_method'],
                                               页内归属依据=item['identity_reason'],检索方式=item.get('retrieval_reason'),匹配详情=item.get('project_matches',[]),
                                               可能所属项目=[targets[p]['name']+'（第'+str(targets[p]['row'])+'行）' for p in projects if p in targets],
                                               原文证据=list(dict.fromkeys(c['evidence'] for c in item['field_candidates'] if c['column']==column)),提示='项目归属待确认，必须核对数值所属项目' if pending else '跨阶段材料识别候选，需人工确认'))
                    method='跨阶段材料：'+candidate['material_stage']+'，'+item['text_method']+('，项目归属待确认' if pending else '，需人工确认')
                    add(f'{get_column_letter(column)}{row["row"]}',candidate['value'],ref,method,False,candidate.get('type'),pending)
    identities={t['project']:t for t in identities.values()}
    for branch in ('6.用地规划','7.工程规划'):
        ledger=read_json(job/'work'/(branch+'_台账.json'),[])
        records=[dict(name=r[12],project_code=r[9],address=r[15],values=r,source_row=i) for i,r in enumerate(ledger[1:],2) if r[12]]
        for row in rows:
            target=identities[row['project']]
            status,ranked=decide(target,records)
            for match in ranked[:5]:
                if match['conflicts']:continue
                record=match['record']['values'];index=match['record']['source_row']
                exact=status=='MATCH_EXACT' and normalize(target.get('name'))==normalize(record[12])
                method=('原始台账：项目代码(Y列)与台账XMDM、项目名称与XMMC双项一致' if exact else
                        '原始台账候选：名称相似度{:.0%}、地址相似度{:.0%}，需人工确认'.format(match['name_score'],match['address_score']))
                for col,source_col in ((25,10),(26,19)):
                    location=f'原始台账!{get_column_letter(source_col)}{index}'
                    ref=source_ref(str(SOURCE/branch/'ZRZYSITE_CONSTRUCT__001012_003b.csv'),location,detail={'说明':'源文件扩展名为 CSV，实际内容为 XLSX','项目名称':record[12],'项目代码':record[9],'建设单位':record[18]})
                    add(f'{get_column_letter(col)}{row["row"]}',record[source_col-1],ref,method,exact)
    conflict_file=job/'outputs/governance/已有值冲突.csv'
    if conflict_file.exists():
        with conflict_file.open(encoding='utf-8-sig') as f:
            for c in csv.DictReader(f):
                if c['cell'] in cells:add(c['cell'],c['candidate'],source_ref(c['source']),c['reason'])
    # 空间结果保留局部、多批次候选供选择，不升级为自动证据。
    spatial_results=read_json(job/'work/spatial_matches.json',[])
    spatial_candidates=[(c['result'],True) for item in spatial_results if item.get('supply_pending') for c in item['candidates']]
    spatial_candidates.extend((item,False) for item in spatial_results if not item.get('supply_pending'))
    for item,pending in spatial_candidates:
        if item.get('resolution',{}).get('excluded'):continue
        row=targets.get(item['project'])
        if not row:continue
        for hit in item['approvals']:
            if hit.get('manual_verdict')=='不一致':continue
            a=hit['attributes'];doc=a.get('PFWH') or a.get('PZWH');date=a.get('PZSJ')
            label=str(SOURCE/'数据图层/3.报批图层/批文报批矢量1016.dbf')
            anchor={'supply':'供地范围','parcels':'首登宗地范围','approvals':'报批批次范围','attachments':'附件矢量范围'}[item.get('anchor_layer','supply')]
            ref=source_ref(label,'OBJECTID='+str(a['OBJECTID']),detail=dict(参考范围=anchor,定位依据=item.get('anchor_method'),匹配方式=hit.get('method','空间相交'),参考范围覆盖率=round(hit['supply_coverage'],6),相交面积平方米=round(hit['intersection_m2'],2),批次=a.get('XMMC') or a.get('PCMC')))
            method='图层候选：'+anchor+'；'+hit.get('method','空间相交')+'（覆盖参考范围{:.1%}，相交{:.0f}平方米），需确认项目归属和适用范围'.format(hit['supply_coverage'],hit['intersection_m2'])
            add(f'AZ{row["row"]}',doc,ref,method,project_pending=pending)
            if date and len(date)==8 and date>'19000101':add(f'BA{row["row"]}',date[:4]+'-'+date[4:6]+'-'+date[6:],ref,method+'；批复日期(PZSJ)',value_type='date',project_pending=pending)
    review=dict(filename=read_json(job/'meta.json',{}).get('filename','目标.xlsx'),headers=HEADERS,rows=rows,cells=cells,sources=sources,decisions={},source_root=str(SOURCE),warnings=[] if master.is_file() else ['未找到可选成果总表，本次按目标表、材料与图层核对'],created_at=dt.datetime.now().isoformat(timespec='seconds'))
    review['project_confirmations']={}
    review['aliases']=read_json(job/'outputs/governance/人工别名.json',[])
    review['source_error_count']=0
    for filename,label in [('material_errors.json','供应扫描件'),('pdf_errors.json','材料解析'),('pdf_render_errors.json','PDF渲染'),('ocr_errors.json','OCR')]:
        count=len(read_json(job/'work'/filename,[]));review['source_error_count']+=count
        if count:review['warnings'].append(label+'有'+str(count)+'项异常，已保留原件并继续其他材料，见来源异常清单')
    for registration in read_json(job/'outputs/governance/图件配准.json',[]):
        path=Path(registration['source_path']).resolve()
        if path.is_file() and any(path.is_relative_to(root.resolve()) for root in (BASE,SOURCE)):
            sources.setdefault(registration['source'],dict(id=registration['source'],label=registration['source_label'],location='配准图件',path=str(path),kind=path.suffix.lower().lstrip('.'),detail={'配准最大误差米':registration['max_error_m']}))
    for cell,decision in previous.get('decisions',{}).items():
        if cell not in cells:continue
        candidate=next((c for c in cells[cell]['candidates'] if c['id']==decision),None)
        if decision=='original' or candidate and (not candidate.get('project_pending') or previous.get('project_confirmations',{}).get(cell)==decision):
            review['decisions'][cell]=decision
            if candidate and candidate.get('project_pending'):review['project_confirmations'][cell]=decision
        else:review['warnings'].append(cell+' 的原候选已失效，请重新核对')
    write_json(job/'review.json',review)
    return review

def spatial_preview(job,project):
    matches=read_json(job/'work/spatial_matches.json',[])
    if project not in {str(r['project']) for r in matches}:raise ValueError('该项目没有空间挂接结果。')
    path=job/'work/spatial_geometry.json'
    if not path.exists():run_step(job,'spatial_matches.py','--preview-only')
    view=read_json(path,{}).get(project)
    if not view:raise ValueError('该项目暂无可绘制的空间图斑。')
    registrations=read_json(job/'outputs/governance/图件配准.json',[])
    if registrations:view['registrations']=[r for r in registrations if r['project']==project and Path(r['source_path']).is_file() and hashlib.sha256(Path(r['source_path']).read_bytes()).hexdigest()==r['source_signature']]
    if (job/'work/coordinate_inputs.json').exists():view['coordinate_inputs']=read_json(job/'work/coordinate_inputs.json',[])
    return view


def resolve_review(review):
    counts=dict(filled=0,conflict=0,candidate=0,selected=0,empty=0,original=0,kept=0)
    for cell,e in review['cells'].items():
        e['value']=e['original'];decision=review.get('decisions',{}).get(cell)
        if decision=='original':e['status']='kept'
        elif decision:
            chosen=next((c for c in e['candidates'] if c['id']==decision),None)
            if chosen:
                if chosen.get('project_pending') and review.get('project_confirmations',{}).get(cell)!=decision:raise ValueError('采用前必须确认该页数值属于当前项目。')
                e['value']=chosen['value'];e['status']='selected'
            else:raise ValueError('选择已过期，请重新检索。')
        else:
            candidates=e['candidates'];different=[c for c in candidates if value_key(c['value'])!=value_key(e['original'])]
            if len(candidates)>1 or (e['original'] is not None and different):e['status']='conflict'
            elif e['original'] is None and len(candidates)==1:
                e['status']='filled' if candidates[0]['auto'] and not candidates[0].get('project_pending') else 'candidate'
                if e['status']=='filled':e['value']=candidates[0]['value']
            else:e['status']='original' if e['original'] is not None else 'empty'
        counts[e['status']]+=1
    review['counts']=counts
    return review


def save_judgement(job,payload):
    review=read_json(job/'review.json');project=payload.get('project');source=payload.get('source')
    if project not in {r['project'] for r in review['rows']}:raise ValueError('项目不存在。')
    if source not in review['sources']:raise ValueError('来源不存在。')
    if payload.get('verdict') not in ('一致','不一致','待核实'):raise ValueError('判读结论无效。')
    note=payload.get('note')
    if not isinstance(note,str) or not 1<=len(note.strip())<=2000:raise ValueError('请填写1至2000字判读依据。')
    view=spatial_preview(job,project)
    views=[c['view'] for c in view['supply_candidates']] if view.get('supply_pending') else [view]
    hits=payload.get('hits',[])
    if not isinstance(hits,list) or not any(payload.get('anchor_layer','supply')==v.get('anchor_layer','supply') and payload.get('supply_indices')==v['supply_indices'] and set(hits).issubset({h['id'] for h in v['hits']}) for v in views):raise ValueError('图斑与当前项目候选不一致，请重新打开空间图。')
    chosen=next(v for v in views if payload.get('anchor_layer','supply')==v.get('anchor_layer','supply') and payload.get('supply_indices')==v['supply_indices'] and set(hits).issubset({h['id'] for h in v['hits']}))
    from relationship_resolution import signature
    entries=read_json(job/'outputs/governance/人工图形判读.json',[])
    entries.append(dict(project=project,source=source,source_label=review['sources'][source]['label'],anchor_layer=payload.get('anchor_layer','supply'),supply_indices=payload['supply_indices'],hits=hits,
                        anchor_signature=signature(chosen.get('supply')),geometry_signatures={h['id']:signature(h.get('geometry')) for h in chosen['hits'] if h['id'] in hits},
                        source_block=(review['sources'][source].get('detail') or {}).get('实体分段'),source_path=review['sources'][source].get('path'),source_signature=hashlib.sha256(Path(review['sources'][source]['path']).read_bytes()).hexdigest() if review['sources'][source].get('path') else None,
                        verdict=payload['verdict'],note=note.strip(),created_at=dt.datetime.now().isoformat(timespec='seconds')))
    write_json(job/'outputs/governance/人工图形判读.json',entries)
    return dict(saved=True,count=len(entries))

def save_alias(job,payload):
    review=read_json(job/'review.json');project=payload.get('project');kind=payload.get('kind');alias=payload.get('alias');canonical=payload.get('canonical')
    if project not in {r['project'] for r in review['rows']} or kind not in ('name','address','road'):raise ValueError('项目或别名类型无效')
    if not isinstance(alias,str) or not 2<=len(alias.strip())<=200:raise ValueError('别名请输入2至200个字符')
    if kind=='road' and (not isinstance(canonical,str) or not 2<=len(canonical.strip())<=100):raise ValueError('请填写道路现名')
    aliases=read_json(job/'outputs/governance/人工别名.json',[])
    aliases=[a for a in aliases if not (a['project']==project and a['kind']==kind and normalize(a['alias'])==normalize(alias))]
    if not payload.get('remove'):aliases.append(dict(project=project,kind=kind,alias=alias.strip(),canonical=canonical.strip() if kind=='road' else None,created_at=dt.datetime.now().isoformat(timespec='seconds')))
    write_json(job/'outputs/governance/人工别名.json',aliases);return dict(saved=True)

def save_relationship_decision(job,payload):
    project=payload.get('project');mode=payload.get('mode');note=payload.get('note')
    if mode not in ('shared','exclusive','reject','reset') or not isinstance(note,str) or not 1<=len(note.strip())<=2000:raise ValueError('请选择范围归属并填写依据')
    view=spatial_preview(job,project);views=[c['view'] for c in view['supply_candidates']] if view.get('supply_pending') else [view]
    from relationship_resolution import signature
    chosen=next((v for v in views if signature(v['supply'])==payload.get('anchor_signature')),None)
    if chosen is None:raise ValueError('参考范围已变化，请重新打开')
    hits=payload.get('hits',[])
    if not isinstance(hits,list) or not set(hits).issubset({h['id'] for h in chosen['hits']}):raise ValueError('所选图斑不存在')
    entries=read_json(job/'outputs/governance/范围归属确认.json',[])
    prior=next((r for r in entries if r['project']==project and r['anchor_signature']==payload['anchor_signature'] and r['mode']==mode),{})
    hashes=dict(prior.get('geometry_signatures',{}));hashes.update({h['id']:signature(h['geometry']) for h in chosen['hits'] if h['id'] in hits})
    entries=[r for r in entries if not (r['project']==project and r['anchor_signature']==payload['anchor_signature'])]
    if mode!='reset':entries.append(dict(project=project,anchor_signature=payload['anchor_signature'],mode=mode,geometry_signatures=hashes,note=note.strip(),created_at=dt.datetime.now().isoformat(timespec='seconds')))
    write_json(job/'outputs/governance/范围归属确认.json',entries);return dict(saved=True)

def save_coordinate_system(job,payload):
    source=payload.get('source');crs=payload.get('crs')
    if source not in read_json(job/'work/coordinate_inputs.json',[]):raise ValueError('来源不在当前矢量清单中')
    if not isinstance(crs,str) or len(crs)>10000:raise ValueError('请填写有效的EPSG代码或WKT坐标系')
    from gis_geometry import runtime
    CRS,_=runtime();value=CRS.from_user_input(crs).to_wkt()
    records=read_json(job/'work/coordinate_systems.json',{});records[source]=value;write_json(job/'work/coordinate_systems.json',records)
    write_json(job/'outputs/governance/人工指定坐标系.json',records);return dict(saved=True)

def save_georeference(job,payload):
    review=read_json(job/'review.json');project=payload.get('project');sid=payload.get('source');source=review['sources'].get(sid,{})
    if project not in {r['project'] for r in review['rows']} or source.get('kind') not in ('jpg','jpeg','png','bmp'):raise ValueError('请选择当前项目及可读取的图片原件；PDF请使用渲染页')
    path=Path(source['path']).resolve()
    if not path.is_file() or not any(path.is_relative_to(root.resolve()) for root in (BASE,SOURCE)):raise ValueError('图片来源无效')
    view=spatial_preview(job,project);views=[c['view'] for c in view['supply_candidates']] if view.get('supply_pending') else [view]
    if not any(v.get('reference_crs')==payload.get('crs') for v in views):raise ValueError('配准坐标系必须与当前参考图层一致')
    from georeference import register_image
    ident=hashlib.sha256((project+'|'+sid).encode()).hexdigest()[:16]
    result=register_image(path,payload.get('points'),payload['crs'],job/'outputs/governance/图件配准'/ident,payload.get('tolerance',2))
    result.update(id=ident,project=project,source=sid,source_path=str(path),source_label=source['label'])
    records=read_json(job/'outputs/governance/图件配准.json',[]);records=[r for r in records if r['id']!=ident];records.append(result)
    write_json(job/'outputs/governance/图件配准.json',records);return result

def save_choice(job,cell,candidate_id,confirm_project=False):
    review=read_json(job/'review.json')
    if cell not in review['cells']:raise ValueError('单元格不存在。')
    entry=review['cells'][cell]
    if candidate_id!='original' and candidate_id is not None and candidate_id not in {c['id'] for c in entry['candidates']}:raise ValueError('候选不存在。')
    chosen=next((c for c in entry['candidates'] if c['id']==candidate_id),None)
    if chosen and chosen.get('project_pending') and confirm_project is not True:raise ValueError('请先查看原件并确认该页数值属于当前项目。')
    review.setdefault('project_confirmations',{}).pop(cell,None)
    if chosen and chosen.get('project_pending'):review['project_confirmations'][cell]=candidate_id
    if candidate_id is None:review.setdefault('decisions',{}).pop(cell,None)
    else:review.setdefault('decisions',{})[cell]=candidate_id
    write_json(job/'review.json',review)
    return resolve_review(review)

def export_review(job):
    review=resolve_review(read_json(job/'review.json'));patches=[];provenance=[]
    for e in review['cells'].values():
        selected=review.get('decisions',{}).get(e['cell'])
        candidate=next((c for c in e['candidates'] if c['id']==selected),None)
        if e['status']=='filled':candidate=e['candidates'][0]
        if candidate and e['status'] in ('filled','selected'):
            patches.append(dict(cell=e['cell'],project=e['project'],value=candidate['value'],type=candidate.get('type'),allow_overwrite=e['status']=='selected',source='；'.join(review['sources'][s]['label']+' '+review['sources'][s]['location'] for s in candidate['sources'])))
            for sid in candidate['sources']:provenance.append(dict(cell=e['cell'],project=e['project'],field=e['field'],value=candidate['value'],status=e['status'],**review['sources'][sid]))
    write_json(job/'work/patches.json',patches)
    write_json(job/'work/cell_states.json',{k:e['status'] for k,e in review['cells'].items()})
    run_step(job,'fill_workbook.mjs');run_step(job,'preserve_original_package.py');run_step(job,'verify_result.py')
    write_json(job/'outputs/governance/用户选择与来源.json',dict(decisions=review.get('decisions',{}),project_confirmations=review.get('project_confirmations',{}),provenance=provenance,counts=review['counts']))
    run_step(job,'governance_outputs.py')
    return job/'outputs/governance/test_result_已填充.xlsx'
