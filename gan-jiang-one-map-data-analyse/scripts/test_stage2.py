import datetime as dt
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from unittest.mock import Mock
from types import SimpleNamespace
from contextlib import closing

import openpyxl
from stage2 import Stage2, BASE, STAGES, read_material, template_fields, digest


class Stage2Tests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=BASE/'work');self.root=Path(self.temp.name)
        self.source=self.root/'source'
        for _,folder in STAGES:(self.source/folder).mkdir(parents=True)
        self.service=Stage2(self.root/'db',self.source);self.service.configure(BASE/'templates/test.xlsx')
        for pid,name in [('a','甲道路工程'),('b','乙道路工程')]:
            self.service.db.execute('INSERT INTO projects VALUES (?,?,?,?,?,1)',(pid,name,'S4',str(self.source/STAGES[3][1]/pid),'{}'))
        self.service.db.commit()

    def tearDown(self):self.service.close();self.temp.cleanup()

    def ingest(self,stage,text):
        path=self.source/STAGES[stage-1][1]/'material.txt';path.write_text(text,encoding='utf-8')
        with self.service.db:self.service.ingest(path,[text])
        return path

    def field(self,code,pid='a'):return next(f for f in self.service.project(pid)['fields'] if f['code']==code)

    def adopt(self,field):
        return self.service.review(dict(project='a',field=field['code'],action='adopt',candidate=field['candidates'][0]['id'],confirm_identity=True))

    def test_s5_contract_fills_s4_and_area_separation(self):
        self.ingest(5,'国有建设用地使用权出让合同；项目名称：甲道路工程；受让人：甲公司；宗地面积：2公顷；合同编号：H1234567890123456；签订日期：2024年3月4日；')
        area=self.field('BI');self.assertEqual(area['candidates'][0]['value'],20000)
        self.assertEqual(self.field('BH')['candidates'][0]['value'],20000)
        self.assertEqual(self.field('BJ')['candidates'][0]['value'],'平方米')
        self.assertIn('2公顷',self.field('BI')['candidates'][0]['raw'])
        self.assertEqual(area['candidates'][0]['stage'],'S5');self.assertEqual(area['candidates'][0]['business_stage'],'土地供应')
        self.adopt(area);self.assertEqual(self.field('BI')['value'],20000)

    def test_s6_conditions_fill_s2(self):
        self.ingest(6,'建设用地规划条件；项目名称：甲道路工程；地块编码：ABC-123；用地面积：123平方米；出具编号：赣规2024001；')
        self.assertEqual(self.field('AM')['candidates'][0]['value'],'ABC-123')
        self.assertEqual(self.field('AN')['candidates'][0]['value'],123)
        self.assertFalse(self.field('BI')['candidates'])

    def test_s7_code_never_merges_parcels(self):
        code='2024-360100-01-01-123456'
        for pid in ('a','b'):self.service.db.execute('UPDATE projects SET original=? WHERE id=?',(json.dumps({'Y':code}),pid))
        self.ingest(7,'项目代码：'+code+'；建设单位：甲公司；')
        self.assertEqual(self.service.db.execute('SELECT count(*) FROM projects').fetchone()[0],2)
        for pid in ('a','b'):
            self.assertEqual(self.field('Y',pid)['value'],code)
            self.assertNotEqual(self.field('Y',pid)['status'],'filled')

    def test_s1_document_in_s4(self):
        self.ingest(4,'建设项目用地预审与选址意见书；项目名称：甲道路工程；拟用地面积：12亩；证书文号：用字第123456789012号；')
        self.assertEqual(self.field('AF')['candidates'][0]['value'],12)
        self.assertFalse(self.field('BI')['candidates'])

    def test_multiple_projects_are_segmented(self):
        self.ingest(5,'国有建设用地使用权出让合同；项目名称：甲道路工程；宗地面积：10平方米；项目名称：乙道路工程；宗地面积：20平方米；')
        self.assertEqual([c['value'] for c in self.field('BI')['candidates']],[10])
        self.assertEqual([c['value'] for c in self.field('BI','b')['candidates']],[20])

    def test_multiple_parcels_conflict_without_overwriting(self):
        self.ingest(5,'国有建设用地使用权出让合同；项目名称：甲道路工程；宗地面积：10平方米；合同编号：H1234567890123456；项目名称：甲道路工程；宗地面积：20平方米；合同编号：H1234567890123457；')
        self.assertEqual(self.field('BI')['status'],'conflict');self.assertIsNone(self.field('BI')['value'])
        self.assertEqual(len(self.field('BO')['candidates']),2)

    def test_original_values_and_review_history_persist(self):
        self.service.db.execute('UPDATE projects SET original=? WHERE id=?',(json.dumps({'BI':99}),'a'))
        path=self.ingest(5,'国有建设用地使用权出让合同；项目名称：甲道路工程；宗地面积：10平方米；')
        self.assertEqual(self.field('BI')['value'],99);self.adopt(self.field('BI'))
        self.service.close();self.service=Stage2(self.root/'db',self.source)
        self.assertEqual(self.field('BI')['value'],10)
        with self.service.db:self.service.ingest(path,[path.read_text(encoding='utf-8')])
        self.assertEqual(self.field('BI')['value'],10)
        self.service.review(dict(project='a',field='BI',action='undo'))
        self.assertEqual(self.field('BI')['value'],99)
        self.assertEqual(len(self.service.project('a')['history']),2)

    def test_changed_source_invalidates_adoption_before_rescan(self):
        path=self.ingest(5,'国有建设用地使用权出让合同；项目名称：甲道路工程；宗地面积：10平方米；')
        self.adopt(self.field('BI'));path.write_text('changed',encoding='utf-8')
        self.assertEqual(self.field('BI')['status'],'stale');self.assertIsNone(self.field('BI')['value'])

    def test_preview_preserves_69_headers_and_dates(self):
        self.service.db.execute("UPDATE projects SET active=0 WHERE id='b'")
        self.ingest(5,'国有建设用地使用权出让合同；项目名称：甲道路工程；签订日期：2024年3月4日；')
        self.adopt(self.field('BP'));path=self.service.preview()
        with closing(openpyxl.load_workbook(path)) as book, closing(openpyxl.load_workbook(BASE/'templates/test.xlsx')) as template:
            self.assertEqual([book['Sheet1'].cell(2,c).value for c in range(1,70)],[template['Sheet1'].cell(2,c).value for c in range(1,70)])
            self.assertEqual(book['Sheet1'].max_column,69)
            self.assertEqual(str(book['Sheet1'].merged_cells),str(template['Sheet1'].merged_cells))
            self.assertEqual(book['Sheet1']['BP3'].value,dt.datetime(2024,3,4))
            self.assertIsNone(book['Sheet1']['X3'].value)
        audit=json.loads((self.service.folder/'preview-audit.json').read_text(encoding='utf-8'))
        self.assertEqual(audit[0]['candidate']['stage'],'S5')

    def test_exclusion_and_undo(self):
        self.ingest(5,'国有建设用地使用权出让合同；项目名称：甲道路工程；宗地面积：10平方米；')
        field=self.field('BI');cid=field['candidates'][0]['id']
        self.service.review(dict(project='a',field='BI',action='exclude',candidate=cid))
        self.assertTrue(self.field('BI')['candidates'][0]['excluded'])
        self.service.review(dict(project='a',field='BI',action='undo'))
        self.assertFalse(self.field('BI')['candidates'][0]['excluded'])

    def test_unassigned_records_can_be_attached_without_filling_all_fields(self):
        self.ingest(5,'国有建设用地使用权出让合同；宗地面积：10平方米；')
        record=self.service.records()[0]
        self.assertEqual(record['projects'],0)
        self.service.attach(dict(project='a',record=record['id'],note='已核对合同原件宗地及位置',confirm_identity=True))
        self.assertIsNone(self.field('BI')['value']);self.assertEqual(len(self.field('BI')['candidates']),1)
        self.assertFalse(self.field('BI','b')['candidates'])

    def test_failure_is_explicit_and_does_not_make_values(self):
        path=self.source/STAGES[4][1]/'bad.pdf';path.write_text('invalid',encoding='utf-8')
        task=self.service.scan();self.assertEqual(task['status'],'partial');self.assertTrue(task['errors'])
        self.assertIsNone(self.field('BI')['value'])

    def test_agent_candidates_require_real_excerpt_and_correct_group(self):
        path=self.ingest(6,'建设用地规划条件；项目名称：甲道路工程；赣新直管〔2020〕038号；')
        record=self.service.records()[0]
        payload=dict(project='a',record=record['id'],field='AQ',raw='赣新直管〔2020〕038号',value='赣新直管〔2020〕038号',reason='已核对文书标题下编号',business_stage='规划条件与红线')
        proposed=self.service.propose(payload)
        self.assertEqual(self.field('AQ')['status'],'candidate')
        self.service.review(dict(project='a',field='AQ',action='adopt',candidate=proposed['candidate'],confirm_identity=True))
        with self.service.db:self.service.ingest(path,[path.read_text(encoding='utf-8')])
        self.assertEqual(self.field('AQ')['value'],payload['value'])
        self.assertEqual(self.field('AQ')['status'],'selected')
        with self.assertRaises(ValueError):self.service.propose(dict(payload,raw='虚构原文'))
        with self.assertRaises(ValueError):self.service.propose(dict(payload,field='BO'))

    def test_only_specific_unique_identifier_can_auto_fill_empty(self):
        self.service.db.execute('UPDATE projects SET original=? WHERE id=?',(json.dumps({'BO':'H1234567890123456'}),'a'))
        self.ingest(5,'国有建设用地使用权出让合同；合同编号：H1234567890123456；宗地面积：10平方米；')
        self.assertEqual(self.field('BI')['status'],'candidate')
        unit='360122009015GB00006W00000000'
        self.service.db.execute('UPDATE projects SET original=? WHERE id=?',(json.dumps({'BO':'H1234567890123456','O':unit}),'a'))
        self.ingest(5,'国有建设用地使用权出让合同；合同编号：H1234567890123456；不动产单元号：'+unit+'；宗地面积：10平方米；')
        self.assertEqual(self.field('BI')['status'],'filled');self.assertEqual(self.field('BI')['value'],10)
        self.service.db.execute('UPDATE projects SET original=? WHERE id=?',(json.dumps({'BO':'H1234567890123456','BI':99}),'a'))
        self.assertEqual(self.field('BI')['value'],99);self.assertEqual(self.field('BI')['status'],'conflict')

    def test_m3_retains_partial_intersection_as_candidate(self):
        from shapely.geometry import box,mapping,shape
        a=self.source/STAGES[3][1]/'a'/'plot.txt';a.parent.mkdir();a.write_text('a')
        b=self.source/STAGES[2][1]/'plot.txt';b.write_text('b')
        with self.service.db:
            self.service.ingest(a,[dict(location='object=1',text='甲道路工程',geometry=mapping(box(0,0,10,10)),crs='EPSG:3857')])
            self.service.ingest(b,[dict(location='object=2',text='乙道路工程',geometry=mapping(box(5,0,15,10)),crs='EPSG:3857')])
        with patch('gis_geometry.projected_crs'),patch('gis_geometry.reproject',side_effect=lambda g,*_:g):
            evidence=self.service.spatial('a')['evidence'][0]
        self.assertEqual(evidence['intersection_m2'],50)
        self.assertEqual(evidence['anchor_coverage'],.5)
        self.assertEqual(evidence['target_coverage'],.5)
        self.assertEqual(self.service.db.execute('SELECT count(*) FROM decisions').fetchone()[0],0)

    def test_shapefile_sidecar_update_changes_material_version(self):
        path=self.source/STAGES[3][1]/'plot.shp';path.write_bytes(b'shape')
        path.with_suffix('.dbf').write_bytes(b'old')
        before=digest(path);path.with_suffix('.dbf').write_bytes(b'new')
        self.assertNotEqual(digest(path),before)

    def test_http_review_route_uses_persistent_service(self):
        import app.server as server
        self.ingest(5,'国有建设用地使用权出让合同；项目名称：甲道路工程；宗地面积：10平方米；')
        cid=self.field('BI')['candidates'][0]['id']
        payload=dict(project='a',field='BI',action='adopt',candidate=cid,confirm_identity=True)
        handler=SimpleNamespace(headers={},path='/api/stage2/review',body=lambda *_:json.dumps(payload).encode(),json=Mock())
        with patch.object(server,'Stage2',side_effect=lambda:Stage2(self.root/'db',self.source)),patch.object(server,'STAGE2_RUNNING',False):
            server.Handler.do_POST(handler)
        result=handler.json.call_args.args[0]
        self.assertNotIn('error',result)
        self.assertEqual(self.field('BI')['value'],10)

    def test_template_semantics_are_not_replaced_by_similar_terms(self):
        self.ingest(3,'建设用地批复；项目名称：甲道路工程；项目类型：交通项目；项目投资：20万元；批次名称：第四批；征收面积：12公顷；耕地面积：3公顷；报批日期：2024年3月4日；是否完成征转：是；')
        self.assertEqual(self.field('AA')['candidates'][0]['value'],'交通项目')
        self.assertEqual(self.field('AW')['candidates'][0]['value'],12)
        self.assertEqual(self.field('AY')['candidates'][0]['value'],'2024-03-04')
        self.assertEqual(self.field('BC')['candidates'][0]['value'],'公顷')
        self.assertEqual(self.field('AX')['candidates'][0]['value'],'是')

    def test_import_preserves_existing_date_and_text_identifier(self):
        path=self.root/'real-input.xlsx'
        with closing(openpyxl.load_workbook(BASE/'templates/test.xlsx')) as book:
            for c in range(1,70):book['Sheet1'].cell(3,c).value=None
            book['Sheet1']['X3']='原始项目';book['Sheet1']['BA3']=dt.datetime(2024,3,4)
            book['Sheet1']['BO3']='0000123456';book.save(path)
        self.assertEqual(self.service.import_values(path),1)
        output=self.service.preview()
        with closing(openpyxl.load_workbook(output)) as book:
            self.assertEqual(book['Sheet1']['BA3'].value,dt.datetime(2024,3,4))
            self.assertEqual(book['Sheet1']['BO3'].value,'0000123456')

    def test_s5_directory_prefix_is_unit_number_not_year_suffix(self):
        unit='360122009006GB00002W00000000'
        for suffix in ('-2009','-','——2026','_2013','无对应供地矢量',''):
            with self.subTest(suffix=suffix):
                folder=self.source/STAGES[4][1]/(unit+suffix);folder.mkdir()
                self.service.db.execute('UPDATE projects SET stage=?,path=? WHERE id=?',('S5',str(folder),'a'))
                target=next(t for t in self.service.targets() if t['project']=='a')
                self.assertEqual(target['unit_id'],unit)
                self.assertIn('unit_id',target['inferred_keys'])
                self.assertEqual(self.service.project('a')['directory_unit_id'],unit)
        self.service.db.execute('UPDATE projects SET original=? WHERE id=?',(json.dumps({'O':'360122009006GB00159W00000000'}),'a'))
        self.assertEqual(next(t for t in self.service.targets() if t['project']=='a')['unit_id'],'360122009006GB00159W00000000')

    def test_s5_directory_hint_does_not_auto_fill_material_fields(self):
        unit='360122009006GB00002W00000000'
        folder=self.source/STAGES[4][1]/(unit+'-2009');folder.mkdir()
        self.service.db.execute('UPDATE projects SET stage=?,path=? WHERE id=?',('S5',str(folder),'a'))
        self.ingest(5,'国有建设用地使用权出让合同；不动产单元号：'+unit+'；宗地面积：10平方米；')
        field=self.field('BI')
        self.assertEqual(field['status'],'candidate');self.assertIsNone(field['value'])
        self.assertEqual(field['candidates'][0]['basis']['directory_unit_id']['value'],unit)


if __name__=='__main__':unittest.main()
