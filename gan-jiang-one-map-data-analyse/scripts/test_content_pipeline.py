import copy
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import io
from types import SimpleNamespace
from material_archives import RarArchive
from PIL import Image
import openpyxl
from pypdf import PdfWriter,PdfReader
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
import shapefile
import extract_pdfs
import extract_fields
import governance_outputs
import vector_outputs
from link_stages import target_records
from review_service import BASE,HEADERS,write_json,read_json,collect_review,resolve_review,save_choice,save_judgement,validate_upload


def workbook(job,names=('华佗路道路工程',)):
    wb=openpyxl.Workbook();sheet=wb.active;sheet.append([None]);sheet.append(HEADERS)
    for name in names:
        row=[None]*69;row[1]='same-archive';row[23]=name;row[17]='3601002020A00636';sheet.append(row)
    wb.save(job/'test_result.xlsx');sheet.delete_rows(3,len(names));wb.save(job/'project_info_all_v5_a - 成果20261021.xlsx');wb.close()


def pdf(path,text='H360202001000012'):
    path.parent.mkdir(parents=True,exist_ok=True)
    writer=PdfWriter();page=writer.add_blank_page(width=400,height=200)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream=DecodedStreamObject();stream.set_data(('BT /F1 12 Tf 20 100 Td ('+text+') Tj ET').encode())
    page[NameObject('/Contents')]=stream;writer.write(path);writer.close()


def evidence(source,stage='土地供应',name='新增道路工程',date='2020-02-03'):
    return dict(source=source,projects=[],pending_projects=[],identity_conflicts=[],material_stage=stage,text='项目名称：'+name+'；',
                fields=dict(contract_id=['H360202001000099'],regulator_id=['3601002020A00999']),
                field_candidates=[dict(column=67,value='H360202001000099',material_stage=stage,evidence='合同编号'),dict(column=68,value=date,material_stage=stage,evidence='签订日期',type='date')],
                identity_reason='尚未关联',text_method='PDF原生文字',storage_stage='6.用地规划')


class ContentPipelineTests(unittest.TestCase):
    def test_missing_archive_number_and_optional_master_do_not_block_review(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);workbook(job);wb=openpyxl.load_workbook(job/'test_result.xlsx');wb.active['B3']=None;wb.save(job/'test_result.xlsx');wb.close()
            (job/'project_info_all_v5_a - 成果20261021.xlsx').unlink()
            validate_upload(job/'test_result.xlsx')
            self.assertEqual(target_records(job/'test_result.xlsx')[0]['archive_id'],'')
            with patch('review_service.SOURCE',job):review=collect_review(job)
            self.assertEqual(len(review['rows']),1);self.assertTrue(review['warnings'])

    def test_registration_projects_retain_ambiguous_units_for_review(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);write_json(job/'work/地表宗地.json',[dict(BDCDYH='unit-1',ZL='位置一'),dict(BDCDYH='shared-unit'),dict(BDCDYH='shared-unit')])
            projects,pending=governance_outputs.registration_projects(job,[],[])
            self.assertEqual([p['identity']['unit_id'] for p in projects],['unit-1']);self.assertEqual(len(pending),2)

    def test_rar_chinese_names_read_only_and_bad_archives_are_reported(self):
        listing=SimpleNamespace(returncode=0,stdout='材料/选址意见书.pdf\n'.encode('gb18030'),stderr=b'')
        with patch('material_archives.shutil.which',return_value='tar'),patch('material_archives.subprocess.run',return_value=listing):
            archive=RarArchive(Path('input.rar'));members=archive.infolist()
        self.assertEqual(members[0].filename,'材料/选址意见书.pdf')
        process=SimpleNamespace(stdout=io.BytesIO(b'PDF bytes'),wait=lambda **kw:0,kill=lambda:None)
        with patch('material_archives.subprocess.Popen',return_value=process) as launch:
            with archive.open(members[0]) as stream:self.assertEqual(stream.read(),b'PDF bytes')
            self.assertIn('--',launch.call_args.args[0])
        with patch('material_archives.subprocess.run',return_value=SimpleNamespace(returncode=1,stderr=b'bad format')):
            with self.assertRaisesRegex(RuntimeError,'RAR 读取失败'):archive.infolist()

    def test_conflicting_or_future_supply_documents_are_not_appended(self):
        item=evidence('a.pdf#page=1')
        conflict=copy.deepcopy(item);conflict['source']='a.pdf#page=2';conflict['field_candidates'].append(dict(column=61,value=100))
        item['field_candidates'].append(dict(column=61,value=200))
        projects,deferred=governance_outputs.rebuild_projects([item,conflict],[])
        self.assertEqual(projects,[]);self.assertEqual(len(deferred),1)
        self.assertEqual(governance_outputs.rebuild_projects([evidence('future.pdf',date='2099-01-01')],[])[0],[])

    def test_unmatched_pdf_and_zip_are_read_and_cached_without_fake_project(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);work=job/'work';work.mkdir();root=job/'sources';root.mkdir();workbook(job)
            target=target_records(job/'test_result.xlsx')[0];target['contract_id']='H360202001000012'
            direct=root/'6.用地规划/unknown/1.pdf';pdf(direct)
            archive=root/'7.工程规划/files.zip';archive.parent.mkdir(parents=True)
            with zipfile.ZipFile(archive,'w') as z:z.writestr('random/2.pdf',direct.read_bytes())
            write_json(work/'files.json',[dict(path=str(direct)),dict(path=str(archive))])
            previous=Path.cwd()
            try:
                os.chdir(job)
                with patch.multiple(extract_pdfs,ROOT=root,OUT=work),patch('extract_pdfs.target_records',return_value=[target]),patch.dict(os.environ,{'GOV_ASSETS':str(work)}):
                    extract_pdfs.main()
                    with patch('extract_pdfs.read_pdf_text',side_effect=AssertionError('cache not reused')):extract_pdfs.main()
                documents=read_json(work/'pdf_documents.json');texts=read_json(work/'pdf_text.json')
                self.assertEqual(len(documents),2);self.assertTrue(all(d['retrieval_projects']==[] for d in documents))
                self.assertTrue(all('H360202001000012' in p['text'] for p in texts))
                with patch.multiple(extract_fields,ROOT=root,OUT=work,DEST=job/'outputs/governance'),patch('extract_fields.target_records',return_value=[target]):extract_fields.main()
                records=[json.loads(s) for s in (job/'outputs/governance/识别字段候选.jsonl').read_text(encoding='utf-8').splitlines()]
                self.assertTrue(all(r['projects']==['row-3'] for r in records))
            finally:os.chdir(previous)

    def test_extensionless_images_enter_ocr_and_keep_source_identity(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);work=job/'work';work.mkdir();root=job/'sources';root.mkdir();workbook(job)
            image=root/'6.用地规划/unknown';image.parent.mkdir();Image.new('RGB',(20,20)).save(image,format='JPEG')
            archive=root/'materials.zip'
            with zipfile.ZipFile(archive,'w') as z:z.writestr('no-extension',image.read_bytes())
            write_json(work/'files.json',[dict(path=str(image)),dict(path=str(archive))])
            previous=Path.cwd()
            try:
                os.chdir(job)
                with patch.multiple(extract_pdfs,ROOT=root,OUT=work),patch('extract_pdfs.target_records',return_value=target_records(job/'test_result.xlsx')),patch.dict(os.environ,{'GOV_ASSETS':str(work)}):extract_pdfs.main()
                manifest=read_json(work/'ocr_manifest.json');self.assertEqual(len(manifest),2)
                self.assertTrue(all(Path(p['path']).is_file() for p in manifest));self.assertTrue(any(p['source']==str(image) for p in manifest))
            finally:os.chdir(previous)

    def test_fuzzy_page_reaches_each_target_and_requires_confirmation(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);work=job/'work';work.mkdir();out=job/'outputs/governance';out.mkdir(parents=True)
            workbook(job,('华佗路道路工程','华佗路道路改建工程'))
            write_json(work/'pdf_text.json',[dict(project='unassigned',retrieval_projects=[],source='random.pdf',page=1,text='建设项目选址意见书；项目名称：华佗路道路改造工程；拟用地面积22亩')])
            write_json(work/'ocr_manifest.json',[])
            with patch.multiple(extract_fields,OUT=work,DEST=out,ROOT=job),patch('extract_fields.target_records',return_value=target_records(job/'test_result.xlsx')):extract_fields.main()
            item=json.loads((out/'识别字段候选.jsonl').read_text(encoding='utf-8'))
            self.assertEqual(set(item['pending_projects']),{'row-3','row-4'})
            self.assertTrue(all('name_score' in m for m in item['project_matches']))
            with patch('review_service.SOURCE',job):review=resolve_review(collect_review(job))
            for cell in ('AF3','AF4'):
                candidate=review['cells'][cell]['candidates'][0]
                self.assertEqual(review['cells'][cell]['status'],'candidate');self.assertTrue(candidate['project_pending'])
                with self.assertRaises(ValueError):save_choice(job,cell,candidate['id'])
                self.assertEqual(save_choice(job,cell,candidate['id'],True)['cells'][cell]['value'],22)

    def test_material_pages_tables_new_projects_and_repeatable_bundle(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);work=job/'work';work.mkdir();out=job/'outputs/governance';out.mkdir(parents=True);workbook(job)
            source=job/'source/unknown/a.pdf';pdf(source)
            source_hash=__import__('hashlib').sha256(source.read_bytes()).hexdigest()
            supply=evidence(str(source)+'#page=1')
            old=evidence('old.pdf#page=1',date='2017-12-31');old['fields']={'contract_id':['H360201701000099']}
            matched=evidence('matched.pdf#page=1');matched['pending_projects']=['row-3']
            selection=dict(supply,source='selection.pdf#page=1',material_stage='预审与选址',projects=['row-3'],fields={'project_code':['2019-360090-48-01-020404']},field_candidates=[dict(column=35,value='用字第123456789012号'),dict(column=32,value=22)])
            write_json(work/'pdf_documents.json',[dict(source=str(source),path=str(source)),dict(source='selection.pdf',path=str(source))])
            (out/'识别字段候选.jsonl').write_text('\n'.join(json.dumps(r) for r in (supply,copy.deepcopy(supply),old,matched,selection)),encoding='utf-8')
            for _ in range(2):
                report=governance_outputs.build(job)
                self.assertEqual(report['new_projects'],1);self.assertEqual(report['selection_records'],1)
                wb=openpyxl.load_workbook(out/'项目表_含补建项目.xlsx');self.assertEqual(wb.active.max_row,4);self.assertEqual(wb.active['X4'].value,'新增道路工程');wb.close()
                with zipfile.ZipFile(out/'治理成果.zip') as z:self.assertIn('选址意见书信息提取.xlsx',z.namelist())
            self.assertEqual(__import__('hashlib').sha256(source.read_bytes()).hexdigest(),source_hash)
            for path in (out/'阶段材料').rglob('*.pdf'):
                with path.open('rb') as stream:self.assertEqual(len(PdfReader(stream).pages),1)
            governance_outputs.table(out/'formula.xlsx',['value'],[['=1+1']])
            wb=openpyxl.load_workbook(out/'formula.xlsx');self.assertEqual(wb.active['A2'].data_type,'s');wb.close()

    def test_attachment_shp_components_zip_and_missing_projects(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);work=job/'work';work.mkdir();workbook(job,('华佗路道路工程','另一个道路工程'))
            stem=job/'source/attachment';stem.parent.mkdir()
            with shapefile.Writer(str(stem)) as writer:
                writer.field('DZJGH','C',30);writer.poly([[(0,0),(0,10),(10,10),(10,0),(0,0)]]);writer.record('3601002020A00636')
            stem.with_suffix('.prj').write_text((__import__('inspect_sources').ROOT/'数据图层/4.供地图层/供地矢量1016.prj').read_text())
            archive=job/'source/a.zip'
            with zipfile.ZipFile(archive,'w') as z:
                for ext in ('.shp','.shx','.dbf','.prj'):z.write(stem.with_suffix(ext),'attachment'+ext)
            write_json(work/'files.json',[dict(path=str(archive))])
            report=vector_outputs.build(job);self.assertEqual(report['attachment_groups'],1);self.assertEqual(report['errors'],[])
            self.assertEqual(len(list((job/'outputs/governance/附件矢量').rglob('*.shp'))),1)
            # Both rows carry the same regulator ID in this fixture, so both are retained.
            self.assertEqual(report['missing_shp_projects'],0)

    def test_visual_judgement_validates_source_geometry_and_persists(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);write_json(job/'review.json',dict(rows=[dict(project='row-3')],sources={'s':dict(label='红线图')}))
            write_json(job/'work/spatial_matches.json',[dict(project='row-3')])
            write_json(job/'work/spatial_geometry.json',{'row-3':dict(supply_indices=[1],hits=[dict(id='approvals:1:0')])})
            payload=dict(project='row-3',source='s',supply_indices=[1],hits=['approvals:1:0'],verdict='待核实',note='图形位置待核对')
            self.assertTrue(save_judgement(job,payload)['saved'])
            with self.assertRaises(ValueError):save_judgement(job,dict(payload,supply_indices=[9]))
            with self.assertRaises(ValueError):save_judgement(job,dict(payload,source='missing'))
            self.assertEqual(len(read_json(job/'outputs/governance/人工图形判读.json')),1)


if __name__=='__main__':unittest.main()
