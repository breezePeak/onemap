import json
import os
import build_patches
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import openpyxl
from link_stages import target_records
from cross_stage_materials import search_terms,path_projects
from review_service import HEADERS,collect_review,export_review
from review_service import BASE,resolve_review,save_choice,validate_upload,write_json
from review_service import reuse_material_cache,read_json

def fixture(original=None,values=(('a','甲',True),)):
    return dict(cells={'Z3':dict(cell='Z3',original=original,value=original,candidates=[dict(id=i,value=v,auto=auto,sources=['source1','source2']) for i,v,auto in values])},decisions={})

class ReviewTests(unittest.TestCase):
    def test_fuzzy_ledger_can_be_selected_exported_and_reused(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);wb=openpyxl.Workbook();s=wb.active;s.append([None]);s.append(HEADERS)
            row=[None]*69;row[1]='duplicate-archive';row[23]='华佗路道路工程';row[42]='condition';row[51]='approval';row[34]='certificate';row[58]='新区';row[60]=15;row[61]='亩';s.append(row)
            wb.save(job/'test_result.xlsx');wb.close()
            identity=target_records(job/'test_result.xlsx')[0]
            self.assertEqual([identity[k] for k in ('condition_id','approval_id','certificate_id')],['condition','approval','certificate'])
            self.assertEqual(identity['area_m2'],10000)
            record=[None]*19;record[12]='华佗路道路建设工程';record[9]='new-code';record[15]='新区';record[18]='建设单位甲'
            write_json(job/'work/6.用地规划_台账.json',[[None]*19,record])
            with patch('review_service.SOURCE',job):review=resolve_review(collect_review(job))
            candidate=review['cells']['Y3']['candidates'][0]
            self.assertFalse(candidate['auto']);self.assertIsNone(review['cells']['Y3']['value'])
            self.assertEqual(review['cells']['Z3']['candidates'][0]['value'],'建设单位甲')
            save_choice(job,'Y3',candidate['id'])
            self.assertEqual(target_records(job/'test_result.xlsx')[0]['project_code'],'new-code')
            with patch('review_service.SOURCE',job):rebuilt=collect_review(job)
            self.assertEqual(rebuilt['decisions']['Y3'],candidate['id'])
            with patch('review_service.run_step'):export_review(job)
            self.assertIn(('Y3','new-code'),[(p['cell'],p['value']) for p in read_json(job/'work/patches.json')])
            source=openpyxl.load_workbook(job/'test_result.xlsx',read_only=True)
            self.assertIsNone(source.active['Y3'].value);source.close()
            out=job/'outputs/governance';out.mkdir(exist_ok=True,parents=True)
            evidence=[dict(projects=['row-3'],fields=dict(contract_id=['C-1'],regulator_id=['R-1'])),dict(projects=['row-3'],fields=dict(contract_id=['C-2'])),dict(projects=[],fields=dict(unit_id=['unconfirmed']))]
            (out/'识别字段候选.jsonl').write_text('\n'.join(json.dumps(e) for e in evidence),encoding='utf-8')
            identity=target_records(job/'test_result.xlsx')[0]
            self.assertEqual(identity['regulator_id'],'R-1');self.assertEqual(identity['inferred_keys'],['regulator_id'])
            self.assertFalse(identity.get('contract_id'));self.assertFalse(identity.get('unit_id'))

    def test_master_archive_is_not_identity_but_address_area_are_candidates(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);wb=openpyxl.Workbook();s=wb.active;s.append([None]);s.append(HEADERS)
            row=[None]*69;row[1]='same';row[23]='华佗路道路工程';row[58]='甲路以东乙路以西';row[60]=1;row[61]='公顷';s.append(row);wb.save(job/'test_result.xlsx')
            row[23]='新建商业办公楼';row[25]='候选建设单位';s.delete_rows(3);s.append(row);wb.save(job/'project_info_all_v5_a - 成果20261021.xlsx');wb.close()
            with patch('review_service.SOURCE',job):review=resolve_review(collect_review(job))
            candidate=review['cells']['Z3']['candidates'][0]
            self.assertFalse(candidate['auto']);self.assertIn('地址相似度100%',candidate['methods'][0]);self.assertIn('面积相对误差0.0%',candidate['methods'][0])
            self.assertIsNone(review['cells']['Z3']['value'])

    def test_duplicate_archive_rows_stay_separate(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);work=job/'work';work.mkdir()
            wb=openpyxl.Workbook();sheet=wb.active;sheet.append([None]);sheet.append(HEADERS)
            for name in ('甲方道路工程','乙方道路工程'):
                row=[None]*69;row[1]='same-001';row[23]=name;sheet.append(row)
            wb.save(job/'test_result.xlsx')
            # Empty master avoids unrelated candidates in this row isolation check.
            sheet.delete_rows(3,2);wb.save(job/'project_info_all_v5_a - 成果20261021.xlsx');wb.close()
            validate_upload(job/'test_result.xlsx')
            targets=target_records(job/'test_result.xlsx')
            self.assertEqual([t['project'] for t in targets],['row-3','row-4'])
            self.assertEqual([t['archive_id'] for t in targets],['same-001']*2)
            terms=search_terms(targets,work,job/'outputs/governance')
            self.assertEqual(path_projects('same-001/合同.pdf',targets,terms),['row-3','row-4'])
            self.assertNotIn('row-3',terms['row-3'])
            ledger=[[None]*19]
            for name,code in (('甲方道路工程','code-a'),('乙方道路工程','code-b')):
                record=[None]*19;record[12]=name;record[9]=code;ledger.append(record)
            write_json(work/'6.用地规划_台账.json',ledger)
            write_json(work/'spatial_matches.json',[])
            previous=Path.cwd()
            try:
                os.chdir(job)
                with patch.multiple(build_patches,OUT=work,DEST=job/'outputs/governance'):build_patches.main()
            finally:os.chdir(previous)
            self.assertEqual([(p['cell'],p['value']) for p in read_json(work/'patches.json')],[('Y3','code-a'),('Y4','code-b')])
            (work/'6.用地规划_台账.json').unlink()
            evidence=[dict(projects=[t['project']],source=t['name']+'.pdf#page=1',text_method='PDF原生文字',
                           identity_reason='页内名称一致',storage_stage='6.用地规划',
                           field_candidates=[dict(column=32,value=value,evidence=str(value),material_stage='预审与选址')])
                      for t,value in zip(targets,(11,22))]
            out=job/'outputs/governance';out.mkdir(parents=True,exist_ok=True)
            (out/'识别字段候选.jsonl').write_text('\n'.join(json.dumps(e) for e in evidence),encoding='utf-8')
            with patch('review_service.SOURCE',job):review=collect_review(job)
            self.assertEqual([r['archive_id'] for r in review['rows']],['same-001']*2)
            self.assertEqual([c['value'] for c in review['cells']['AF3']['candidates']],[11])
            self.assertEqual([c['value'] for c in review['cells']['AF4']['candidates']],[22])
            save_choice(job,'AF3',review['cells']['AF3']['candidates'][0]['id'])
            selected=save_choice(job,'AF4',review['cells']['AF4']['candidates'][0]['id'])
            self.assertEqual(selected['cells']['AF3']['value'],11)
            self.assertEqual(selected['cells']['AF4']['value'],22)
            with patch('review_service.run_step'):export_review(job)
            self.assertEqual([(p['cell'],p['value']) for p in read_json(work/'patches.json')],[('Y3','code-a'),('AF3',11),('Y4','code-b'),('AF4',22)])

    def test_reuse_successes_across_jobs_but_not_failures_or_choices(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            root=Path(folder);assets=root/'work';old=assets/'jobs/old/work';new=assets/'jobs/new'
            pdf=root/'source.pdf';pdf.write_bytes(b'pdf');stat=pdf.stat()
            good=dict(size=stat.st_size,mtime_ns=stat.st_mtime_ns,pages=['文字'])
            bad=root/'bad.pdf';bad.write_bytes(b'bad')
            write_json(old/'pdf_text_cache.json',{str(pdf):good,str(bad):dict(good,error='failed')})
            image=root/'p1.png';image.write_bytes(b'image')
            record=dict(source='source.pdf',size=3,mtimeMs=1,complete=True,pages=[dict(path=str(image))])
            write_json(old/'render_manifests/good.json',record)
            write_json(old/'render_manifests/bad.json',dict(record,complete=False,error='failed'))
            write_json(old.parent/'review.json',dict(decisions={'Z3':'old-choice'}))
            reuse_material_cache(new,assets)
            self.assertEqual(read_json(new/'work/pdf_text_cache.json'),{str(pdf):good})
            self.assertEqual(read_json(new/'work/render_manifests/good.json'),record)
            self.assertFalse((new/'work/render_manifests/bad.json').exists())
            self.assertFalse((new/'review.json').exists())
            pdf.write_bytes(b'changed PDF')
            another=assets/'jobs/another';reuse_material_cache(another,assets)
            self.assertEqual(read_json(another/'work/pdf_text_cache.json'),{})

    def test_unique_evidence_fills_blank(self):
        r=resolve_review(fixture());self.assertEqual(r['cells']['Z3']['status'],'filled');self.assertEqual(r['cells']['Z3']['value'],'甲')
        self.assertEqual(len(r['cells']['Z3']['candidates'][0]['sources']),2)
    def test_conflict_never_overwrites_original(self):
        r=resolve_review(fixture('原值'));self.assertEqual(r['cells']['Z3']['status'],'conflict');self.assertEqual(r['cells']['Z3']['value'],'原值')
    def test_multiple_candidates_never_auto_fill(self):
        r=resolve_review(fixture(values=(('a','甲',True),('b','乙',True))));self.assertEqual(r['cells']['Z3']['status'],'conflict');self.assertIsNone(r['cells']['Z3']['value'])
    def test_weak_candidate_requires_user(self):
        r=resolve_review(fixture(values=(('a','甲',False),)));self.assertEqual(r['cells']['Z3']['status'],'candidate');self.assertIsNone(r['cells']['Z3']['value'])
    def test_project_pending_requires_explicit_confirmation_and_persists(self):
        r=fixture();r['cells']['Z3']['candidates'][0]['project_pending']=True
        self.assertEqual(resolve_review(r)['cells']['Z3']['status'],'candidate')
        with tempfile.TemporaryDirectory(dir=BASE/'work') as d:
            job=Path(d);write_json(job/'review.json',r)
            with self.assertRaises(ValueError):save_choice(job,'Z3','a')
            with self.assertRaises(ValueError):save_choice(job,'Z3','a','true')
            chosen=save_choice(job,'Z3','a',True)
            self.assertEqual(chosen['cells']['Z3']['status'],'selected');self.assertEqual(chosen['project_confirmations'],{'Z3':'a'})
            self.assertEqual(save_choice(job,'Z3',None)['project_confirmations'],{})
        r['decisions']={'Z3':'a'}
        with self.assertRaises(ValueError):resolve_review(r)
    def test_choice_persists_and_can_be_undone(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as d:
            assert Path(d).resolve().is_relative_to((BASE/'work').resolve())
            job=Path(d);write_json(job/'review.json',fixture('原值'))
            self.assertEqual(save_choice(job,'Z3','a')['cells']['Z3']['value'],'甲')
            self.assertEqual(save_choice(job,'Z3','original')['cells']['Z3']['status'],'kept')
            self.assertEqual(save_choice(job,'Z3',None)['cells']['Z3']['status'],'conflict')
            with self.assertRaises(ValueError):save_choice(job,'Z3','missing')
    def test_non_xlsx_rejected(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as d:
            assert Path(d).resolve().is_relative_to((BASE/'work').resolve())
            path=Path(d)/'bad.xlsx';path.write_text('not an excel file')
            with self.assertRaises(ValueError):validate_upload(path)

if __name__=='__main__':unittest.main()
