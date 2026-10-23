import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch,Mock
from types import SimpleNamespace
import zipfile
from PIL import Image
from image_materials import prepare_image,valid_image
from review_service import BASE,HEADERS,collect_review,write_json,read_json
import extract_materials
import extract_pdfs
import governance_outputs
import openpyxl

class BadMaterialTests(unittest.TestCase):
    def test_failed_job_can_retry_but_running_job_cannot_be_enqueued_twice(self):
        import app.server as server
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            job=Path(folder);write_json(job/'meta.json',dict(status='failed',error='old failure'))
            handler=SimpleNamespace(headers={},path='/api/jobs/0123456789abcdef/rerun',json=Mock())
            with patch('app.server.job_path',return_value=job),patch.object(server.POOL,'submit') as submit:
                server.Handler.do_POST(handler)
                submit.assert_called_once_with(server.execute,job)
                self.assertEqual(read_json(job/'meta.json')['status'],'queued');self.assertIsNone(read_json(job/'meta.json')['error'])
                server.Handler.do_POST(handler);self.assertEqual(submit.call_count,1);self.assertEqual(handler.json.call_args.args[1],400)

    def fixture(self,root):
        work=root/'work';work.mkdir();archive=root/'4.土地供应阶段/划拨.zip';archive.parent.mkdir()
        buffer=io.BytesIO();Image.new('RGB',(20,20),'white').save(buffer,format='JPEG')
        files={'划拨/2023-02628 测试/image00007.jpg':b'','划拨/2023-02628 测试/corrupt.jpg':b'not an image','划拨/2023-02628 测试/valid.jpg':buffer.getvalue()}
        with zipfile.ZipFile(archive,'w') as z:
            for name,value in files.items():z.writestr(name,value)
        write_json(work/'划拨_zip.json',[dict(name=n,raw=n,size=len(v)) for n,v in files.items()])
        return work,archive,files

    def test_empty_and_corrupt_members_do_not_block_valid_or_poison_cache(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            root=Path(folder);work,archive,files=self.fixture(root);targets=[dict(project='row-3',archive_id='2023-02628')]
            assets=root/'assets';dest=assets/'materials/row-3';dest.mkdir(parents=True)
            valid_name=next(n for n in files if n.endswith('valid.jpg'));cached=dest/(hashlib.sha256(valid_name.encode()).hexdigest()[:12]+'_valid.jpg');cached.write_bytes(b'')
            os.utime(cached,ns=(archive.stat().st_mtime_ns+1_000_000_000,)*2)
            with patch.multiple(extract_materials,ROOT=root,OUT=work,RecordedZip=zipfile.ZipFile),patch('extract_materials.target_records',return_value=targets),patch.dict(os.environ,GOV_ASSETS=str(assets)):
                extract_materials.main()
            manifest=read_json(work/'ocr_manifest.json');errors=read_json(work/'material_errors.json')
            self.assertEqual(len(manifest),1);self.assertTrue(valid_image(manifest[0]['path']));self.assertEqual(len(errors),2)
            self.assertTrue(any('0字节' in e['error'] for e in errors));self.assertTrue(any('corrupt.jpg' in e['source'] for e in errors))
            with zipfile.ZipFile(archive) as z:self.assertEqual(z.read(next(iter(files))),b'')
            self.assertFalse(list(assets.rglob('*.writing')))

    def test_failed_replacement_preserves_old_image_but_is_not_current_evidence(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            dest=Path(folder)/'cached.jpg';prepare_image(io.BytesIO(self.jpeg()),dest);before=dest.read_bytes()
            with self.assertRaises(ValueError):prepare_image(io.BytesIO(b''),dest)
            self.assertEqual(dest.read_bytes(),before);self.assertFalse(list(Path(folder).glob('*.writing')))

    def jpeg(self):
        buffer=io.BytesIO();Image.new('RGB',(20,20)).save(buffer,format='JPEG');return buffer.getvalue()

    def test_full_text_ingestion_excludes_empty_image_and_continues(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            root=Path(folder);work=root/'work';work.mkdir();empty=root/'empty.jpg';empty.write_bytes(b'');good=root/'valid.jpg';good.write_bytes(self.jpeg())
            selected=[dict(path=p,kind='.jpg',projects=[],source=str(p),storage_stage='4.土地供应阶段') for p in (empty,good)]
            previous=Path.cwd()
            try:
                os.chdir(root)
                with patch.multiple(extract_pdfs,ROOT=root,OUT=work),patch('extract_pdfs.target_records',return_value=[]),patch('extract_pdfs.discover_files',return_value=(selected,{})),patch.dict(os.environ,GOV_ASSETS=str(work)):extract_pdfs.main()
            finally:os.chdir(previous)
            manifest=read_json(work/'ocr_manifest.json');self.assertEqual(len(manifest),1);self.assertEqual(manifest[0]['source'],str(good))
            self.assertIn('0字节',read_json(work/'pdf_errors.json')[0]['error'])

    def test_split_entities_survive_governance_export_and_errors_are_visible(self):
        with tempfile.TemporaryDirectory(dir=BASE/'work') as folder:
            root=Path(folder);work=root/'work';work.mkdir();out=root/'outputs/governance';out.mkdir(parents=True)
            wb=openpyxl.Workbook();wb.active.append([None]);wb.active.append(HEADERS)
            for name in ('甲方道路工程','乙方道路工程'):
                row=[None]*69;row[23]=name;wb.active.append(row)
            wb.save(root/'test_result.xlsx');wb.close()
            image=root/'page.jpg';image.write_bytes(self.jpeg());records=[]
            for index,(project,name,area) in enumerate([('row-3','甲方道路工程',10),('row-4','乙方道路工程',20)],1):
                records.append(dict(source='same.jpg',block_id='block-'+str(index),projects=[project],pending_projects=[],identity_conflicts=[],image=str(image),text='项目名称：'+name+'；',material_stage='预审与选址',fields={},field_candidates=[dict(column=32,value=area,material_stage='预审与选址',evidence=str(area))],identity_reason='实体内名称一致',text_method='本机OCR'))
            (out/'识别字段候选.jsonl').write_text('\n'.join(json.dumps(r) for r in records),encoding='utf-8')
            write_json(work/'material_errors.json',[dict(source='empty.jpg',error='图片为空文件（0字节）')])
            with patch('review_service.SOURCE',root):review=collect_review(root)
            self.assertEqual(review['source_error_count'],1);self.assertTrue(any('扫描件' in w for w in review['warnings']))
            report=governance_outputs.build(root)
            self.assertEqual(report['recognized_pages'],1);self.assertEqual(report['recognized_entities'],2);self.assertEqual(report['selection_records'],2);self.assertEqual(report['source_errors'],1)
            wb=openpyxl.load_workbook(out/'选址意见书信息提取.xlsx',read_only=True)
            self.assertEqual([r[4] for r in wb.active.iter_rows(min_row=2,values_only=True)],['10','20']);wb.close()
            self.assertTrue((out/'来源异常清单.xlsx').is_file())

if __name__=='__main__':unittest.main()
