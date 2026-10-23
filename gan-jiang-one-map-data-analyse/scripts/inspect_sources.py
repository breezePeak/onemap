"""只读盘点：识别伪扩展名台账、ZIP 目录与 DBF 属性。"""
import json
import struct
import zipfile
import sqlite3
from contextlib import closing
from functools import lru_cache
from pathlib import Path
import openpyxl
from extract_archives import RecordedZip,OUTPUT

from source_config import SOURCE as ROOT
OUT = Path('work'); OUT.mkdir(exist_ok=True)


def is_office_lock(path):
    return Path(path).name.startswith('~$')


def extracted_files(database,root,active=None):
    if not database.is_file():return []
    root=Path(root).resolve()
    @lru_cache(maxsize=None)
    def canonical(path): return str(Path(path).resolve())
    with closing(sqlite3.connect(database.resolve().as_uri()+'?mode=ro',uri=True)) as db:
        rows=db.execute("SELECT archive,raw,dest,dsize,dmtime FROM files WHERE status='done'").fetchall()
        names={}
        if db.execute("SELECT 1 FROM sqlite_master WHERE name='archives'").fetchone():
            for archive,members in db.execute('SELECT path,members FROM archives'):
                names.update({(archive,m['raw']):m['name'] for m in json.loads(members)})
    if active is not None: rows=[r for r in rows if (canonical(r[0]),r[1]) in active]
    parents={canonical(dest):(archive,raw) for archive,raw,dest,_,_ in rows if dest and Path(dest).suffix.lower() in ('.zip','.rar','.7z')}
    def label(archive,raw,seen=()):
        key=canonical(archive)
        name=names.get((key,raw),raw).replace('\\','/')
        if key in seen:raise ValueError('解压索引存在循环引用')
        if key in parents:
            parent,member=parents[key];return label(parent,member,(*seen,key))+'!'+name
        return str(Path(archive))+'!'+name
    result=[]
    for archive,raw,dest,size,mtime in rows:
        path=Path(dest)
        if not path.is_file() or (active is not None and not Path(canonical(str(path))).is_relative_to(root)):continue
        stat=path.stat()
        if (stat.st_size,stat.st_mtime_ns)!=(size,mtime):continue
        source=label(archive,raw);origin=Path(source.split('!',1)[0])
        if not origin.is_relative_to(root):continue
        result.append(dict(path=str(path),size=stat.st_size,source=source,storage_stage=origin.relative_to(root).parts[0]))
    return result

def dbf(path):
    data = path.read_bytes()
    count, header, length = struct.unpack_from('<IHH', data, 4)
    fields = []
    offset = 32
    while data[offset] != 13:
        raw_name = data[offset:offset+11].split(b'\0')[0]
        try: name = raw_name.decode('utf-8')
        except UnicodeError: name = raw_name.decode('gb18030')
        fields.append((name, chr(data[offset+11]), data[offset+16]))
        offset += 32
    records = []
    for i in range(count):
        start = header + i * length
        if data[start:start+1] == b'*': continue
        pos = start + 1; record = {}
        for name, kind, size in fields:
            raw = data[pos:pos+size]; pos += size
            try: value = raw.decode('utf-8')
            except UnicodeError: value = raw.decode('gb18030', errors='replace')
            record[name] = value.strip(' \0')
        records.append(record)
    return fields, records

def main():
    inventory = []
    for path in ROOT.rglob('*'):
        if not is_office_lock(path) and path.is_file():
            inventory.append({'path': str(path), 'size': path.stat().st_size})
    existing={str(Path(x['path']).resolve()) for x in inventory}
    inventory.extend(x for x in extracted_files(OUTPUT/'解压记录.sqlite',ROOT) if str(Path(x['path']).resolve()) not in existing)
    (OUT/'files.json').write_text(json.dumps(inventory, ensure_ascii=False), encoding='utf-8')
    for path in (ROOT/'数据图层').rglob('*.dbf'):
        if '房屋' in str(path): continue
        fields, records = dbf(path)
        print(path.name, len(records), fields)
        print(records[:1])
        (OUT/(path.stem+'.json')).write_text(json.dumps(records,ensure_ascii=False),encoding='utf-8')
    for path in (ROOT/'4.土地供应阶段').glob('*.zip'):
        with RecordedZip(path) as z:
            rows=[]
            for info in z.infolist():
                name=info.filename
                if not info.flag_bits & 0x800:
                    try:name=name.encode('cp437').decode('gb18030')
                    except (UnicodeError,LookupError):pass
                rows.append({'name':name,'raw':info.filename,'size':info.file_size})
            (OUT/(path.stem+'_zip.json')).write_text(json.dumps(rows,ensure_ascii=False),encoding='utf-8')
            print('ZIP',path.name,len(rows));print(rows[:5])
    for path in ROOT.glob('*/*.csv'):
        if is_office_lock(path):continue
        if path.read_bytes()[:2] != b'PK':continue
        with path.open('rb') as f:
            w=openpyxl.load_workbook(f,read_only=True,data_only=True)
            for s in w:
                rows=list(s.iter_rows(max_col=60,values_only=True))
                (OUT/(path.parent.name+'_台账.json')).write_text(json.dumps(rows,ensure_ascii=False,default=str),encoding='utf-8')
                print('伪CSV台账',path.parent.name,s.title,s.max_row,s.max_column,rows[:2])

if __name__=='__main__':main()
