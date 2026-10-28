"""已目视核对原件的真实语义样本；程序性采用/撤销仅作用于验证库。"""
from contextlib import closing
import datetime as dt
import json
from pathlib import Path
import openpyxl

from stage2 import Stage2,BASE,SOURCE,read_material,digest


def main():
    store=BASE/'work/stage2-real-validation';results=[]
    # 仅测试样本固定期望值；生产提取逻辑不按这些文件名或项目硬编码。
    cases=[('S5','江西紫祥光电照明有限公司国有建设用地使用权出让合同.PDF',2,'BO','3620130104010K','3620130104010K','土地供应'),
           ('S6','038号86亩中小学用地规划条件.pdf',1,'AQ','赣新直管〔2020〕038号','赣新直管〔2020〕038号','规划条件与红线')]
    with closing(Stage2(store,SOURCE)) as service:
        for stage,name,page,code,raw,value,business in cases:
            path=next(p for p in SOURCE.rglob(name) if p.is_file() and p.name==name)
            pages,method=read_material(path,store/'cache'/digest(path))
            with service.db:mid=service.ingest(path,pages,method)
            targets=[t for t in service.targets() if Path(t['path']).is_dir() and path.is_relative_to(Path(t['path']))]
            assert len(targets)==1,'样本目标目录不唯一，需要重新选择真实样本'
            pid=targets[0]['project']
            record=service.db.execute('SELECT id FROM records WHERE material=? AND location=?',(mid,f'page={page};block-1')).fetchone()[0]
            proposed=service.propose(dict(project=pid,record=record,field=code,raw=raw,value=value,business_stage=business,
                reason='开发验证样本：已目视核对原件对应页；文书与当前目录记录关联仍须业务审核，不推广为自动填写规则'))
            service.review(dict(project=pid,field=code,action='adopt',candidate=proposed['candidate'],confirm_identity=True,note='仅验证库的程序性审核测试，非用户业务确认'))
            field=next(f for f in service.project(pid)['fields'] if f['code']==code)
            assert field['value']==value and field['status']=='selected'
            with closing(Stage2(store,SOURCE)) as reopened:
                field=next(f for f in reopened.project(pid)['fields'] if f['code']==code)
                assert field['value']==value,'重开数据库后采用结果丢失'
            preview=service.preview()
            with closing(openpyxl.load_workbook(preview,read_only=True,data_only=True)) as book:
                assert book['Sheet1'].max_column==69
                assert value in [r[0] for r in book['Sheet1'].iter_rows(min_row=3,min_col=field['column'],max_col=field['column'],values_only=True)]
            audit=json.loads((store/'preview-audit.json').read_text(encoding='utf-8'))
            assert any(a['candidate']['record']==record and a['candidate']['stage']==stage for a in audit)
            service.review(dict(project=pid,field=code,action='undo',note='验证结束，恢复候选待审核'))
            results.append(dict(storage_stage=stage,target_field=code,business_stage=business,value=value,path=str(path),page=page,fingerprint=digest(path),
                checks=['真实原文片段校验','业务分组约束','候选采用','数据库重开持久化','69列预览写入','独立来源审计','撤销恢复候选'],
                caveat='仅验证库程序性审核，不代表生产项目归属已由用户确认'))
        report=dict(created=dt.datetime.now().isoformat(),cases=results,production_reviews_changed=False)
        (BASE/'docs/stage2-semantic-real-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))


if __name__=='__main__':main()
